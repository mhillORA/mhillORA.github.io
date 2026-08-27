"""
Ingest Dry Eye Overview.xlsx → Cosmos (ora-clinical-recruiting / crcscheduling)

Same containers as anterior-segment ingest:
  legacy-studies, legacy-sites, legacy-study-site-outcomes

Workbook shape differs (per-study tabs with a mid-sheet site table), but docs are
the same schema. Studies/sites that already exist are matched 1:1 by name and
REUSE existing ids — no duplicate study/site docs. New ones are created.

Outcomes use source=dry-eye-overview. If an outcome for the same
studyId|siteId|group already exists (e.g. from anterior), funnel fields are
preserved and we only annotate sourceFiles.

Usage:
  python ingest/legacy_dry_eye_overview.py --dry-run
  python ingest/legacy_dry_eye_overview.py --apply
  python ingest/legacy_dry_eye_overview.py path/to/file.xlsx --apply
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

try:
    from azure.cosmos import CosmosClient, PartitionKey
except ImportError:
    print("Install: pip install azure-cosmos openpyxl")
    sys.exit(1)

from openpyxl import load_workbook

DEFAULT_XLSX = Path(
    r"c:\Users\shue1\Downloads\Dry Eye Overview-Moved to Sharepoint, this is no longer active.xlsx"
)
ENDPOINT = os.environ.get("COSMOS_ENDPOINT", "https://ora-clinical-recruiting.documents.azure.com:443/")
KEY = os.environ.get("COSMOS_KEY", "")
DATABASE_ID = os.environ.get("DATABASE_ID", "crcscheduling")
STUDIES_CONTAINER = "legacy-studies"
SITES_CONTAINER = "legacy-sites"
OUTCOMES_CONTAINER = "legacy-study-site-outcomes"
SOURCE = "dry-eye-overview"
INDICATION = "Dry Eye"

SKIP_SHEETS = {
    "pre-screenings",
    "sheet38",
    "aldeyra overview",
}

# Rows that look like site labels but aren't
SKIP_SITE_NAMES = {
    "sites",
    "site",
    "site #",
    "site number",
    "total",
    "totals",
    "average",
    "averages",
    "confirmed",
    "waiting for site",
    "need to send to site",
    "previous sites",
    "row labels",
    "visit 1",
    "day 1",
    "confirmed dates",
}

GROUP_RE = re.compile(
    r"^(?P<base>.+?)\s*\(?\s*g(?:roup)?\s*(?P<g>\d+)\s*\)?\s*$",
    re.I,
)


def load_key_from_local_settings():
    global KEY, ENDPOINT, DATABASE_ID
    if KEY:
        return
    repo = Path(__file__).resolve().parents[1]
    candidates = [
        repo / "data-api-connections.json",
        repo / "api" / "local.settings.json",
        Path(r"C:\Users\shue1\OneDrive\Desktop\HoldAll\CHAOS\azure-api-fixed\local.settings.json"),
    ]
    for p in candidates:
        if not p.exists():
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        vals = data.get("Values") or {}
        KEY = vals.get("COSMOS_KEY") or KEY
        ENDPOINT = vals.get("COSMOS_ENDPOINT") or ENDPOINT
        DATABASE_ID = vals.get("DATABASE_ID") or DATABASE_ID
        conn = (data.get("cosmosdb-connection") or {}).get("connectionString") or ""
        if conn and not KEY:
            parts = {}
            for chunk in conn.rstrip(";").split(";"):
                if "=" in chunk:
                    k, v = chunk.split("=", 1)
                    parts[k.strip()] = v.strip()
            KEY = parts.get("AccountKey") or KEY
            ENDPOINT = parts.get("AccountEndpoint") or ENDPOINT
            DATABASE_ID = (data.get("cosmosdb-connection") or {}).get("database") or DATABASE_ID
        if KEY:
            return


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return float(v)
    t = str(v).strip().replace(",", "")
    if not t or t == "-" or t.startswith("#"):
        return None
    try:
        return float(t)
    except ValueError:
        m = re.match(r"^([+-]?\d+(?:\.\d+)?)", t)
        if m:
            try:
                return float(m.group(1))
            except ValueError:
                return None
        return None


def s(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v).strip()
    if t.startswith("="):
        return None
    return t or None


def slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (name or "").strip().lower()).strip("-")
    return (slug or "unknown")[:80]


def norm_name(name: str) -> str:
    t = (name or "").strip().lower().replace("&", " and ")
    t = re.sub(r"\b(inc\.?|llc|corp\.?|ltd\.?|pllc|pc|md|dr\.?|sc)\b", " ", t, flags=re.I)
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def outcome_id(study: str, site: str, group) -> str:
    raw = f"{study}|{site}|{group}"
    h = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]
    return f"legacy-outcome-{h}"


def site_id_for(name: str) -> str:
    return f"legacy-site-{slugify(name)}"


def study_id_for(name: str) -> str:
    return f"legacy-study-{slugify(name)}"


def parse_group_from_site(site_label: str) -> tuple[str, float | None]:
    """'Torkildsen G2' → ('Torkildsen', 2.0); strip (Doubles) noise."""
    label = (site_label or "").strip()
    label = re.sub(r"\s*\(?\s*doubles\s*\)?\s*$", "", label, flags=re.I).strip()
    m = GROUP_RE.match(label)
    if m:
        return m.group("base").strip(), float(m.group("g"))
    return label, None


def sheet_study_fallback(sheet_name: str) -> str:
    name = re.sub(r"-?\s*complete\s*$", "", sheet_name or "", flags=re.I).strip()
    return name or sheet_name


def col_aliases(header: str) -> str | None:
    h = norm_name(header)
    if not h:
        return None
    if h in {"sites", "site"} or h.startswith("site name"):
        return "site"
    if h in {"site", "site number"} or h in {"site number", "site no"} or h.startswith("site "):
        # site # / site number are not the name column
        if "number" in h or h in {"site"}:
            if h == "sites" or h == "site name":
                return "site"
            if "num" in h or h == "site":
                # ambiguous: only treat pure "sites" as name; site # is index
                if h in {"site number", "site no"} or h.endswith("number") or h == "site":
                    # "site" alone sometimes means name when there's no "sites" col —
                    # handled by prefer_sites below
                    return "site_num" if ("num" in h or h.endswith("#")) else "site_maybe"
        if h == "sites":
            return "site"
    if h == "sites":
        return "site"
    if "start" in h and "date" in h:
        return "start"
    if ("end" in h and "date" in h) or h in {"end date ip", "end ip"}:
        return "end"
    if h == "status":
        return "status"
    if "target" in h and "schedul" in h:
        return "target_scheduled"
    if "target" in h and "enroll" in h:
        return "target_enrolled"
    if h in {"number scheduled", "number scheduled v1", "scheduled v1", "v1 spots"} or (
        "schedul" in h and "v1" in h and "target" not in h
    ):
        return "scheduled"
    if h == "number scheduled" or (h.startswith("number scheduled") and "v2" not in h and "v3" not in h and "v0" not in h and "v4" not in h):
        return "scheduled"
    if "screen" in h and "target" not in h and "to" not in h:
        return "screened"
    if ("actual" in h and "enroll" in h) or h in {"number enrolled", "enrolled"} or (
        "enroll" in h and "target" not in h and "6 week" not in h and "1 year" not in h and "long" not in h
    ):
        return "enrolled"
    if "comment" in h:
        return "comments"
    return None


def map_headers(cells: list) -> dict[str, int] | None:
    """Return role→col index if this looks like a site table header."""
    roles: dict[str, int] = {}
    raw = [s(c) or "" for c in cells]
    norms = [norm_name(x) for x in raw]

    for i, h in enumerate(raw):
        n = norms[i]
        if not n:
            continue
        if n == "sites" or n == "site name":
            roles["site"] = i
        elif n in {"site number", "site no"} or n.replace(" ", "") in {"site#", "sitenumber"}:
            roles["site_num"] = i
        elif n == "site":
            roles.setdefault("site_maybe", i)
        elif "start" in n and "date" in n:
            roles["start"] = i
        elif ("end" in n and "date" in n) or n in {"end date ip", "end ip"}:
            roles["end"] = i
        elif n == "status":
            roles["status"] = i
        elif "target" in n and "schedul" in n:
            roles["target_scheduled"] = i
        elif "target" in n and "enroll" in n:
            roles["target_enrolled"] = i
        elif n in {"number scheduled", "number scheduled v1", "scheduled v1"} or (
            "schedul" in n and "v1" in n and "target" not in n and "v2" not in n
        ):
            roles.setdefault("scheduled", i)
        elif n == "number scheduled" or (
            n.startswith("number scheduled")
            and not any(x in n for x in ("v2", "v3", "v4", "v0"))
        ):
            roles.setdefault("scheduled", i)
        elif "screen" in n and "target" not in n and " to " not in f" {n} ":
            roles.setdefault("screened", i)
        elif n in {"actual enrolled", "number enrolled", "enrolled"} or (
            "enroll" in n and "target" not in n and "total" not in n and "remaining" not in n
        ):
            # Prefer plain Actual Enrolled; else first enroll col (e.g. 6-Week)
            if n in {"actual enrolled", "number enrolled", "enrolled"}:
                roles["enrolled"] = i
            else:
                roles.setdefault("enrolled", i)

    if "site" not in roles and "site_maybe" in roles:
        roles["site"] = roles.pop("site_maybe")
    elif "site_maybe" in roles:
        roles.pop("site_maybe", None)

    if "site" not in roles:
        return None
    # Real site tables have dates and/or targets — reject visit schedules & rate tables
    has_dates = "start" in roles or "end" in roles
    has_target = "target_enrolled" in roles or "target_scheduled" in roles
    has_funnel = any(k in roles for k in ("scheduled", "screened", "enrolled"))
    if not has_funnel:
        return None
    if not (has_dates or has_target):
        return None
    # Header cell for site must literally be Sites / Site Name (not a random "Sites" label mid-row)
    site_hdr = norm_name(raw[roles["site"]] if roles["site"] < len(raw) else "")
    if site_hdr not in {"sites", "site name"}:
        return None
    return roles


def extract_study_name(rows: list) -> str | None:
    """Find 'Study Name' header and value on the next data row."""
    for i, row in enumerate(rows[:25]):
        cells = [s(c) for c in (row or [])]
        for j, c in enumerate(cells):
            if c and c.strip().lower() == "study name":
                # value usually next row, same column
                if i + 1 < len(rows):
                    nxt = rows[i + 1] or []
                    if j < len(nxt) and s(nxt[j]):
                        return s(nxt[j])
                # sometimes value is to the right
                for k in range(j + 1, len(cells)):
                    if cells[k]:
                        return cells[k]
    return None


def resolve_study_label(sheet_name: str, extracted: str | None) -> str:
    """Unique-enough study label; keep protocol tokens from the sheet tab."""
    base = (extracted or "").strip() or sheet_study_fallback(sheet_name)
    fb = sheet_study_fallback(sheet_name)
    proto = re.search(
        r"(ADX-\d+|CYS-\d+|VELOS-\d+|Saturn[-\s]?\d+|OPP\s*\d+|RGN[-\s]?\d+)",
        sheet_name,
        re.I,
    )
    if proto:
        token = re.sub(r"\s+", "-", proto.group(1).strip())
        if norm_name(token) not in norm_name(base):
            return f"{base} ({token})"
    # Sheet tab often more specific than a short Study Name cell ("Aldeyra")
    if len(norm_name(fb)) >= len(norm_name(base)) + 4:
        return fb
    return base


def is_stop_row(row, roles: dict) -> bool:
    vals = [s(c) for c in (row or []) if s(c)]
    if not vals:
        return False
    joined = " ".join(vals).lower()
    if "confirmed dates" in joined or "visit schedule" in joined:
        return True
    site_i = roles["site"]
    site_val = s(row[site_i]) if row and len(row) > site_i else None
    if site_val and site_val.lower() in SKIP_SITE_NAMES:
        return True
    if site_val and site_val.lower().startswith("visit "):
        return True
    return False


def parse_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    records = []
    sheet_stats = []

    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower() in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 3:
            continue

        study_name = resolve_study_label(sheet_name, extract_study_name(rows))
        # Collect every site-table header (PART 1 / PART 2, etc.)
        headers: list[tuple[int, dict]] = []
        for i, row in enumerate(rows[:120]):
            mapped = map_headers(list(row or []))
            if mapped:
                headers.append((i, mapped))
        if not headers:
            sheet_stats.append({"sheet": sheet_name, "study": study_name, "rows": 0, "skip": "no_header"})
            continue

        n = 0
        for hi, (header_i, roles) in enumerate(headers):
            end_at = headers[hi + 1][0] if hi + 1 < len(headers) else min(len(rows), header_i + 80)
            blanks = 0
            for row in rows[header_i + 1 : end_at]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    blanks += 1
                    if blanks >= 3:
                        break
                    continue
                blanks = 0
                if is_stop_row(row, roles):
                    # PART labels / section breaks — keep scanning within window
                    site_i = roles["site"]
                    sv = s(row[site_i]) if len(row) > site_i else None
                    if sv and sv.lower() in {"part 1", "part 2", "part 3"}:
                        continue
                    break

                site_i = roles["site"]
                site_raw = s(row[site_i]) if len(row) > site_i else None
                if not site_raw:
                    continue
                if site_raw.lower() in SKIP_SITE_NAMES:
                    continue
                if re.fullmatch(r"\d+(\.0+)?", site_raw):
                    continue
                if site_raw.lower().startswith("part "):
                    continue

                site_name, group = parse_group_from_site(site_raw)

                def cell(role, _roles=roles, _row=row):
                    if role not in _roles:
                        return None
                    idx = _roles[role]
                    return _row[idx] if len(_row) > idx else None

                scheduled = num(cell("scheduled"))
                screened = num(cell("screened"))
                enrolled = num(cell("enrolled"))
                target_scheduled = num(cell("target_scheduled"))
                target_enrolled = num(cell("target_enrolled"))
                if target_scheduled is None and target_enrolled is not None:
                    target_scheduled = target_enrolled

                start = s(cell("start"))
                end = s(cell("end"))
                if scheduled is None and screened is None and enrolled is None and start is None:
                    continue

                records.append(
                    {
                        "sheet": sheet_name,
                        "study": study_name,
                        "site": site_name,
                        "site_raw": site_raw,
                        "group": group,
                        "pi": site_name,
                        "visit1_start": start,
                        "lplv": end,
                        "target_scheduled": target_scheduled,
                        "scheduled": scheduled,
                        "screened": screened,
                        "enrolled": enrolled,
                    }
                )
                n += 1

        sheet_stats.append({"sheet": sheet_name, "study": study_name, "rows": n, "skip": None})

    wb.close()
    return records, sheet_stats


def match_study(name: str, existing_studies: list[dict]) -> dict | None:
    """1:1 match to an existing legacy study — prefer protocol codes (ADX-019, Saturn-2)."""
    if not name:
        return None

    by_norm: dict[str, dict] = {}
    for st in existing_studies:
        for key in filter(None, [st.get("name"), st.get("title")]):
            kn = norm_name(key)
            if kn and kn not in by_norm:
                by_norm[kn] = st
            kc = kn.replace(" ", "")
            if kc and kc not in by_norm:
                by_norm[kc] = st

    candidates = [name]
    stripped = re.sub(r"\s*\([^)]+\)\s*$", "", name).strip()
    if stripped and stripped != name:
        candidates.append(stripped)

    # Protocol / short codes embedded in dry-eye labels
    for src in (name, stripped):
        for m in re.finditer(
            r"\b(ADX-\d+|CYS-\d+|VELOS-\d+|Saturn[-\s]?\d+|OPP\s*\d+|RGN[-\s]?\d+|BRIM|HanAll|Allergan|Vanda|Palatin|Aerie|Aurinia|Nicox|Kowa|Mitotech)\b",
            src or "",
            re.I,
        ):
            candidates.append(m.group(1))
            # Saturn-2 style already in Cosmos as Saturn-2
            if re.match(r"Saturn", m.group(1), re.I):
                candidates.append(re.sub(r"\s+", "-", m.group(1)))

    # Exact normalized hits first (highest confidence)
    for cand in candidates:
        n = norm_name(cand)
        if not n:
            continue
        if n in by_norm:
            return by_norm[n]
        nc = n.replace(" ", "")
        if nc in by_norm:
            return by_norm[nc]

    # Prefix/startswith only when candidate is a short protocol-like code (>=5) or equal family
    best = None
    best_score = 0
    for cand in candidates:
        n = norm_name(cand)
        if len(n) < 5:
            continue
        for kn, st in by_norm.items():
            if len(kn) < 4:
                continue
            score = 0
            if kn == n:
                score = 1000
            elif kn.replace(" ", "") == n.replace(" ", ""):
                score = 900
            # ADX-019 exact family: existing short name equals token
            elif len(n) <= 12 and (kn == n or kn.startswith(n) or n.startswith(kn)):
                score = 100 + min(len(kn), len(n))
            if score > best_score:
                best, best_score = st, score
    if best_score >= 100:
        return best
    return None


def site_source_weight(site: dict | None) -> int:
    """Prefer feasibility sites so dry-eye outcomes line up with survey data."""
    if not site:
        return 0
    src = (site.get("source") or "").lower()
    sources = site.get("sources") or []
    if not isinstance(sources, list):
        sources = [sources] if sources else []
    blob = " ".join([src, *[str(x).lower() for x in sources]])
    if "feasibility" in blob:
        return 50
    if "anterior" in blob:
        return 10
    if SOURCE in blob:
        return 1
    return 5


def build_pi_site_index(existing_outcomes: list[dict], existing_sites: list[dict]):
    """Map normalized PI surname → preferred legacy site id/name.

    Votes prefer **feasibility** sites so Legacy Sites surveys line up.
    Dry-eye stub sites (source=dry-eye-overview only) are never chosen as targets.
    """
    from collections import Counter

    site_by_id = {s["id"]: s for s in existing_sites}
    votes: dict[str, Counter] = defaultdict(Counter)

    def is_dry_eye_stub(site: dict | None) -> bool:
        if not site:
            return False
        return (site.get("source") or "") == SOURCE

    for o in existing_outcomes:
        pi = o.get("pi")
        sid = o.get("siteId")
        if not pi or not sid:
            continue
        site = site_by_id.get(sid)
        if is_dry_eye_stub(site):
            continue
        key = norm_name(re.sub(r"\s*\([^)]*\)\s*$", "", str(pi)))
        if len(key) < 3:
            continue
        w = 2 + site_source_weight(site)
        votes[key][sid] += w

    for site in existing_sites:
        if is_dry_eye_stub(site):
            continue
        w = site_source_weight(site)
        for field in ("pi", "name"):
            raw = site.get(field)
            if not raw:
                continue
            n = norm_name(raw)
            tokens = n.split()
            for tok in tokens:
                if len(tok) >= 5:
                    votes[tok][site["id"]] += w
            if len(tokens) >= 2 and len(tokens[-1]) >= 4:
                votes[tokens[-1]][site["id"]] += w

    # Known typo / alias → canonical surname key
    aliases = {
        "mclaurn": "mclaurin",
        "elharazi": "el harazi",
        "el-harazi": "el harazi",
        "nieman": "neiman",
    }

    index: dict[str, tuple[str, str]] = {}
    for key, ctr in votes.items():
        # If any feasibility candidate exists for this PI, prefer those only
        # (so dry-eye outcomes line up with site-survey-* data)
        feas_sids = {
            sid
            for sid in ctr
            if site_source_weight(site_by_id.get(sid)) >= 50
        }
        pool = {sid: ctr[sid] for sid in feas_sids} if feas_sids else dict(ctr)
        ranked = sorted(
            pool.items(),
            key=lambda kv: (kv[1], site_source_weight(site_by_id.get(kv[0]))),
            reverse=True,
        )
        sid = ranked[0][0]
        site = site_by_id.get(sid) or {}
        index[key] = (sid, site.get("name") or sid)

    for alias, canon in aliases.items():
        if alias not in index and canon in index:
            index[alias] = index[canon]
    return index


def site_label_keys(label: str) -> list[str]:
    """Candidate surname/keys from a dry-eye site cell."""
    base, _ = parse_group_from_site(label)
    keys = [base]
    # "Jerkins (Loden)" / "Jerkins Site #2" → also try leading surname
    head = re.sub(r"\s*[\(#].*$", "", base or "").strip()
    head = re.sub(r"\s+site\b.*$", "", head, flags=re.I).strip()
    if head and head.lower() not in {k.lower() for k in keys}:
        keys.append(head)
    for m in re.finditer(r"\(([^)]+)\)", label or ""):
        inner = m.group(1).strip()
        if inner and not re.fullmatch(r"doubles|s|g\d+", inner, re.I):
            # "PI: C. Spearman, MD" → Spearman
            inner2 = re.sub(r"^(PI:?\s*)", "", inner, flags=re.I)
            keys.append(inner2)
            toks = re.findall(r"[A-Za-z]{4,}", inner2)
            keys.extend(toks)
    if "/" in (base or ""):
        keys.extend(p.strip() for p in base.split("/") if p.strip())
    # Drop noisy keys
    out = []
    for k in keys:
        k = re.sub(r"\s+", " ", (k or "").strip())
        if not k or k.lower().startswith("site "):
            continue
        if re.fullmatch(r"\d+", k):
            continue
        out.append(k)
    return out


def match_site(
    name: str,
    existing_sites: list[dict],
    pi_index: dict[str, tuple[str, str]] | None = None,
) -> dict | None:
    if not name:
        return None
    pi_index = pi_index or {}

    exact: dict[str, dict] = {}
    for site in existing_sites:
        sn = site.get("name") or ""
        kn = norm_name(sn)
        if kn and kn not in exact:
            exact[kn] = site

    # 1) Exact name
    n = norm_name(name)
    if n in exact:
        return exact[n]
    nc = n.replace(" ", "")
    for kn, site in exact.items():
        if kn.replace(" ", "") == nc:
            return site

    # 2) PI surname index from existing outcomes (true 1:1 when PI known)
    for key in site_label_keys(name):
        kn = norm_name(key)
        if kn in pi_index:
            sid, sname = pi_index[kn]
            # synthesize doc-shaped dict
            hit = next((s for s in existing_sites if s["id"] == sid), None)
            if hit:
                return hit
            return {"id": sid, "name": sname}
        # also try last token of multi-word key
        toks = kn.split()
        if len(toks) >= 2 and toks[-1] in pi_index:
            sid, sname = pi_index[toks[-1]]
            hit = next((s for s in existing_sites if s["id"] == sid), None)
            if hit:
                return hit
            return {"id": sid, "name": sname}

    # 3) Whole-token surname in site name (min 5 chars) — no partial "Peter"/"Peterson"
    best = None
    best_score = 0
    for key in site_label_keys(name):
        kn = norm_name(key)
        if len(kn) < 5:
            continue
        for site in existing_sites:
            sn = norm_name(site.get("name") or "")
            tokens = set(sn.split())
            if kn in tokens:
                score = 50 + len(kn)
                if score > best_score:
                    best, best_score = site, score
            # "Jerkins (S)" style already exact-handled; "Bergstrom Eye" contains bergstrom
            elif len(kn) >= 6 and any(t == kn or t.startswith(kn + " ") for t in [sn]):
                if kn in sn.split() or sn.startswith(kn + " "):
                    score = 40 + len(kn)
                    if score > best_score:
                        best, best_score = site, score
    if best_score >= 45:
        return best
    return None


def build_docs(records, existing_studies, existing_sites, existing_outcomes, source_file: str):
    now = datetime.utcnow().isoformat() + "Z"
    match_report = {
        "studies_matched": [],
        "studies_new": [],
        "sites_matched": [],
        "sites_new": [],
        "outcomes_new": 0,
        "outcomes_existing_preserved": 0,
    }

    pi_index = build_pi_site_index(existing_outcomes, existing_sites)
    # Never match onto dry-eye stubs — prefer feasibility / anterior / other
    matchable_sites = [s for s in existing_sites if (s.get("source") or "") != SOURCE]

    # Cache matches
    study_resolve: dict[str, tuple[str, str, bool]] = {}  # name → (id, canonical_name, is_new)
    site_resolve: dict[str, tuple[str, str, bool]] = {}

    for r in records:
        study = r["study"]
        if study not in study_resolve:
            hit = match_study(study, existing_studies)
            if hit:
                study_resolve[study] = (hit["id"], hit.get("name") or study, False)
                match_report["studies_matched"].append(
                    {"from": study, "to": hit.get("name"), "id": hit["id"]}
                )
            else:
                study_resolve[study] = (study_id_for(study), study, True)
                match_report["studies_new"].append({"name": study, "id": study_id_for(study)})

        site_name = r["site"]
        if site_name not in site_resolve:
            hit = match_site(site_name, matchable_sites, pi_index)
            if hit:
                site_resolve[site_name] = (hit["id"], hit.get("name") or site_name, False)
                match_report["sites_matched"].append(
                    {"from": site_name, "to": hit.get("name"), "id": hit["id"]}
                )
            else:
                site_resolve[site_name] = (site_id_for(site_name), site_name, True)
                match_report["sites_new"].append({"name": site_name, "id": site_id_for(site_name)})

    # Dedupe match report lists
    def uniq(items, key):
        seen = set()
        out = []
        for it in items:
            k = it.get(key)
            if k in seen:
                continue
            seen.add(k)
            out.append(it)
        return out

    match_report["studies_matched"] = uniq(match_report["studies_matched"], "id")
    match_report["studies_new"] = uniq(match_report["studies_new"], "id")
    match_report["sites_matched"] = uniq(match_report["sites_matched"], "id")
    match_report["sites_new"] = uniq(match_report["sites_new"], "id")

    by_study = defaultdict(
        lambda: {
            "target_scheduled": 0.0,
            "scheduled": 0.0,
            "screened": 0.0,
            "enrolled": 0.0,
            "sites": set(),
            "pis": set(),
            "dates": [],
            "lplvs": [],
            "study_id": None,
            "canonical": None,
            "is_new": True,
        }
    )
    by_site = defaultdict(
        lambda: {
            "name": None,
            "target_scheduled": 0.0,
            "scheduled": 0.0,
            "screened": 0.0,
            "enrolled": 0.0,
            "studies": set(),
            "pis": set(),
            "nOutcomeRows": 0,
            "is_new": True,
        }
    )

    outcomes = []
    for r in records:
        study_id, study_canon, study_new = study_resolve[r["study"]]
        site_id, site_canon, site_new = site_resolve[r["site"]]
        oid = outcome_id(study_canon, site_canon, r["group"])
        outcomes.append(
            {
                "id": oid,
                "type": "legacyStudySiteOutcome",
                "studyId": study_id,
                "studyName": study_canon,
                "siteId": site_id,
                "siteName": site_canon,
                "group": r["group"],
                "pi": r["pi"],
                "visit1Start": r["visit1_start"],
                "lplv": r["lplv"],
                "targetScheduled": r["target_scheduled"],
                "scheduled": r["scheduled"],
                "screened": r["screened"],
                "enrolled": r["enrolled"],
                "uniqueId": None,
                "source": SOURCE,
                "sourceFile": source_file,
                "workbookSiteLabel": r["site_raw"],
                "workbookStudyLabel": r["study"],
                "workbookSheet": r["sheet"],
                "ingestedAt": now,
                "createdAt": now,
                "updatedAt": now,
                "_studyIsNew": study_new,
                "_siteIsNew": site_new,
            }
        )

        agg = by_study[study_id]
        agg["study_id"] = study_id
        agg["canonical"] = study_canon
        agg["is_new"] = study_new
        for f_src, f_dst in [
            ("target_scheduled", "target_scheduled"),
            ("scheduled", "scheduled"),
            ("screened", "screened"),
            ("enrolled", "enrolled"),
        ]:
            agg[f_dst] += r.get(f_src) or 0
        agg["sites"].add(site_canon)
        if r["pi"]:
            agg["pis"].add(r["pi"])
        if r["visit1_start"]:
            agg["dates"].append(r["visit1_start"])
        if r["lplv"]:
            agg["lplvs"].append(r["lplv"])

        sagg = by_site[site_id]
        sagg["name"] = site_canon
        sagg["is_new"] = site_new
        sagg["nOutcomeRows"] += 1
        sagg["studies"].add(study_canon)
        if r["pi"]:
            sagg["pis"].add(r["pi"])
        for f_src, f_dst in [
            ("target_scheduled", "target_scheduled"),
            ("scheduled", "scheduled"),
            ("screened", "screened"),
            ("enrolled", "enrolled"),
        ]:
            sagg[f_dst] += r.get(f_src) or 0

    # Only build NEW study/site docs (matched ones keep Cosmos identity + metrics)
    studies = []
    for study_id, agg in by_study.items():
        if not agg["is_new"]:
            continue
        studies.append(
            {
                "id": study_id,
                "type": "legacyStudy",
                "name": agg["canonical"],
                "title": agg["canonical"],
                "therapeuticArea": INDICATION,
                "indication": INDICATION,
                "oraProjectNumber": None,
                "sponsor": None,
                "phase": None,
                "status": "Completed",
                "notes": None,
                "source": SOURCE,
                "sourceFile": source_file,
                "editableFields": True,
                "metrics": {
                    "targetScheduled": round(agg["target_scheduled"], 2),
                    "scheduled": round(agg["scheduled"], 2),
                    "screened": round(agg["screened"], 2),
                    "enrolled": round(agg["enrolled"], 2),
                    "nSites": len(agg["sites"]),
                    "nPis": len(agg["pis"]),
                    "nSiteRows": len([o for o in outcomes if o["studyId"] == study_id]),
                    "visit1StartMin": min(agg["dates"]) if agg["dates"] else None,
                    "visit1StartMax": max(agg["dates"]) if agg["dates"] else None,
                    "lplvMin": min(agg["lplvs"]) if agg["lplvs"] else None,
                    "lplvMax": max(agg["lplvs"]) if agg["lplvs"] else None,
                },
                "createdAt": now,
                "updatedAt": now,
                "ingestedAt": now,
            }
        )

    sites = []
    for sid, agg in sorted(by_site.items(), key=lambda x: -(x[1]["enrolled"])):
        if not agg["is_new"]:
            continue
        sites.append(
            {
                "id": sid,
                "type": "legacySite",
                "name": agg["name"],
                "siteCode": slugify(agg["name"]).upper().replace("-", "_")[:32],
                "status": "Active",
                "notes": None,
                "relationshipPreference": None,
                "advantages": None,
                "disadvantages": None,
                "relationshipNotes": None,
                "linkedArtemisSiteId": None,
                "source": SOURCE,
                "sourceFile": source_file,
                "editableFields": True,
                "indication": INDICATION,
                "indicationsCovered": [INDICATION],
                "therapeuticAreas": [INDICATION],
                "metrics": {
                    "targetScheduled": round(agg["target_scheduled"], 2),
                    "scheduled": round(agg["scheduled"], 2),
                    "screened": round(agg["screened"], 2),
                    "enrolled": round(agg["enrolled"], 2),
                    "nStudies": len(agg["studies"]),
                    "nPis": len(agg["pis"]),
                    "nOutcomeRows": agg["nOutcomeRows"],
                    "studyNames": sorted(agg["studies"]),
                },
                "createdAt": now,
                "updatedAt": now,
                "ingestedAt": now,
            }
        )

    return studies, sites, outcomes, match_report


def fetch_existing(db):
    studies_c = db.get_container_client(STUDIES_CONTAINER)
    sites_c = db.get_container_client(SITES_CONTAINER)
    outcomes_c = db.get_container_client(OUTCOMES_CONTAINER)
    studies = list(
        studies_c.query_items(
            query="SELECT c.id, c.name, c.title, c.indication, c.source FROM c",
            enable_cross_partition_query=True,
        )
    )
    sites = list(
        sites_c.query_items(
            query="SELECT c.id, c.name, c.source, c.linkedArtemisSiteId FROM c",
            enable_cross_partition_query=True,
        )
    )
    outcomes = list(
        outcomes_c.query_items(
            query=(
                "SELECT c.id, c.studyId, c.siteId, c.siteName, c.pi, c[\"group\"] AS grp, "
                "c.scheduled, c.screened, c.enrolled, c.source, c.sourceFile, c.createdAt FROM c"
            ),
            enable_cross_partition_query=True,
        )
    )
    for o in outcomes:
        if "grp" in o:
            o["group"] = o.pop("grp")
    return studies, sites, outcomes


def upsert_all(studies, sites, outcomes, match_report, dry_run=False):
    load_key_from_local_settings()
    if not KEY:
        raise SystemExit("COSMOS_KEY missing. Set env COSMOS_KEY or use local.settings.json")

    client = CosmosClient(ENDPOINT, credential=KEY)
    db = client.get_database_client(DATABASE_ID)
    existing_studies, existing_sites, existing_outcomes = fetch_existing(db)
    by_outcome_id = {o["id"]: o for o in existing_outcomes}

    print("\n=== MATCH REPORT ===")
    print(f"Existing Cosmos: {len(existing_studies)} studies, {len(existing_sites)} sites, {len(existing_outcomes)} outcomes")
    print(f"Studies matched 1:1: {len(match_report['studies_matched'])}")
    for m in match_report["studies_matched"][:25]:
        print(f"  study  {m['from']!r} -> {m['to']!r} ({m['id']})")
    if len(match_report["studies_matched"]) > 25:
        print(f"  ... +{len(match_report['studies_matched']) - 25} more")
    print(f"Studies NEW: {len(match_report['studies_new'])}")
    for m in match_report["studies_new"]:
        print(f"  + {m['name']!r} ({m['id']})")
    print(f"Sites matched 1:1: {len(match_report['sites_matched'])}")
    for m in match_report["sites_matched"][:30]:
        print(f"  site   {m['from']!r} -> {m['to']!r} ({m['id']})")
    if len(match_report["sites_matched"]) > 30:
        print(f"  ... +{len(match_report['sites_matched']) - 30} more")
    print(f"Sites NEW: {len(match_report['sites_new'])}")
    for m in match_report["sites_new"][:40]:
        print(f"  + {m['name']!r} ({m['id']})")
    if len(match_report["sites_new"]) > 40:
        print(f"  ... +{len(match_report['sites_new']) - 40} more")

    to_insert = []
    to_preserve = []
    for doc in outcomes:
        clean = {k: v for k, v in doc.items() if not k.startswith("_")}
        prev = by_outcome_id.get(clean["id"])
        if prev:
            # Preserve existing funnel entirely; annotate provenance only.
            # Never stamp workbook labels onto anterior rows.
            merged = dict(prev)
            files = merged.get("sourceFiles") or (
                [] if not merged.get("sourceFile") else [merged.get("sourceFile")]
            )
            if not isinstance(files, list):
                files = [files] if files else []
            if clean.get("sourceFile") and clean["sourceFile"] not in files:
                files.append(clean["sourceFile"])
            prev_sources = merged.get("sources") or []
            if not isinstance(prev_sources, list):
                prev_sources = [prev_sources] if prev_sources else []
            if merged.get("source") and merged["source"] not in prev_sources:
                prev_sources.append(merged["source"])
            if SOURCE not in prev_sources:
                prev_sources.append(SOURCE)
            merged["sourceFiles"] = files
            merged["sources"] = [x for x in dict.fromkeys(prev_sources) if x]
            merged["updatedAt"] = clean["updatedAt"]
            # Fill blanks only — never overwrite anterior numbers
            for k in ("scheduled", "screened", "enrolled", "targetScheduled", "visit1Start", "lplv", "pi"):
                if merged.get(k) in (None, "") and clean.get(k) not in (None, ""):
                    merged[k] = clean[k]
            to_preserve.append(merged)
            match_report["outcomes_existing_preserved"] += 1
        else:
            to_insert.append(clean)
            match_report["outcomes_new"] += 1

    print(f"\nOutcomes NEW: {match_report['outcomes_new']}")
    print(f"Outcomes existing (preserve funnel, annotate): {match_report['outcomes_existing_preserved']}")
    print(f"New study docs to create: {len(studies)}")
    print(f"New site docs to create: {len(sites)}")

    if dry_run:
        print("\nDRY RUN — no writes.")
        if outcomes:
            print("sample new outcome:", json.dumps(
                {k: v for k, v in outcomes[0].items() if not k.startswith("_")}, indent=2
            )[:800])
        return match_report

    pk_id = PartitionKey(**{"path": "/id"})
    pk_study = PartitionKey(**{"path": "/studyId"})
    db.create_container_if_not_exists(id=STUDIES_CONTAINER, partition_key=pk_id)
    db.create_container_if_not_exists(id=SITES_CONTAINER, partition_key=pk_id)
    db.create_container_if_not_exists(id=OUTCOMES_CONTAINER, partition_key=pk_study)
    studies_c = db.get_container_client(STUDIES_CONTAINER)
    sites_c = db.get_container_client(SITES_CONTAINER)
    outcomes_c = db.get_container_client(OUTCOMES_CONTAINER)

    # Soft-link new sites to live ARTEMIS sites
    live_by_name = {}
    live_by_norm = {}

    def norm_site_name(x: str) -> str:
        t = (x or "").lower()
        t = re.sub(r"\b(inc\.?|llc|corp\.?|ltd\.?|pllc|pc|md|dr\.?|sc)\b", "", t, flags=re.I)
        t = re.sub(r"[^a-z0-9]+", " ", t)
        return " ".join(t.split())

    try:
        live_c = db.get_container_client("sites")
        for item in live_c.query_items(
            query="SELECT c.id, c.name FROM c",
            enable_cross_partition_query=True,
        ):
            if item.get("name"):
                live_by_name[str(item["name"]).strip().lower()] = item["id"]
                nn = norm_site_name(item["name"])
                if nn and nn not in live_by_norm:
                    live_by_norm[nn] = item["id"]
    except Exception as e:
        print("Note: could not read live sites for linking:", e)

    print(f"Creating {len(studies)} new legacy studies...")
    for doc in studies:
        studies_c.upsert_item(doc)

    # For matched studies: ensure Dry Eye on indication if empty; tag sourceFiles
    matched_ids = {m["id"] for m in match_report["studies_matched"]}
    if matched_ids:
        print(f"Annotating {len(matched_ids)} matched studies with dry-eye source...")
        for st in existing_studies:
            if st["id"] not in matched_ids:
                continue
            # Annotate ONLY — never overwrite name/metrics/edits/feasibility fields
            full = studies_c.read_item(st["id"], partition_key=st["id"])
            changed = False
            files = full.get("sourceFiles") or ([] if not full.get("sourceFile") else [full["sourceFile"]])
            if not isinstance(files, list):
                files = [files]
            src_path = outcomes[0]["sourceFile"] if outcomes else DEFAULT_XLSX.name
            if str(src_path) not in files:
                files.append(str(src_path))
                full["sourceFiles"] = files
                changed = True
            sources = full.get("sources") or []
            if not isinstance(sources, list):
                sources = [sources] if sources else []
            if full.get("source") and full["source"] not in sources:
                sources.append(full["source"])
            if SOURCE not in sources:
                sources.append(SOURCE)
                full["sources"] = sources
                changed = True
            # Fill indication only when blank — never replace an existing value
            if not full.get("indication"):
                full["indication"] = INDICATION
                if not full.get("therapeuticArea"):
                    full["therapeuticArea"] = INDICATION
                changed = True
            if changed:
                full["updatedAt"] = datetime.utcnow().isoformat() + "Z"
                studies_c.upsert_item(full)

    print(f"Creating {len(sites)} new legacy sites...")
    for doc in sites:
        if not doc.get("linkedArtemisSiteId"):
            name = (doc.get("name") or "").strip()
            link = live_by_name.get(name.lower()) or live_by_norm.get(norm_site_name(name))
            if link:
                doc["linkedArtemisSiteId"] = link
        sites_c.upsert_item(doc)

    matched_site_ids = {m["id"] for m in match_report["sites_matched"]}
    if matched_site_ids:
        print(f"Annotating {len(matched_site_ids)} matched sites with dry-eye indication/source...")
        for site in existing_sites:
            if site["id"] not in matched_site_ids:
                continue
            # Annotate ONLY — never overwrite name, metrics, relationship, or survey links
            full = sites_c.read_item(site["id"], partition_key=site["id"])
            changed = False
            inds = full.get("indicationsCovered") or []
            if not isinstance(inds, list):
                inds = [inds] if inds else []
            if INDICATION not in inds:
                inds.append(INDICATION)
                full["indicationsCovered"] = inds
                changed = True
            tas = full.get("therapeuticAreas") or []
            if not isinstance(tas, list):
                tas = [tas] if tas else []
            if INDICATION not in tas:
                tas.append(INDICATION)
                full["therapeuticAreas"] = tas
                changed = True
            sources = full.get("sources") or []
            if not isinstance(sources, list):
                sources = [sources] if sources else []
            if full.get("source") and full["source"] not in sources:
                sources.append(full["source"])
            if SOURCE not in sources:
                sources.append(SOURCE)
                full["sources"] = sources
                changed = True
            if changed:
                full["updatedAt"] = datetime.utcnow().isoformat() + "Z"
                sites_c.upsert_item(full)

    print(f"Upserting {len(to_insert)} new outcomes + {len(to_preserve)} annotated existing...")
    batch = 0
    for doc in to_insert + to_preserve:
        outcomes_c.upsert_item(doc)
        batch += 1
        if batch % 100 == 0:
            print(f"  ...{batch}/{len(to_insert) + len(to_preserve)}")
    print("Done.")
    return match_report


def is_dry_eye_outcome(o: dict) -> bool:
    """Only true dry-eye rows — do not treat anterior rows that were annotated."""
    if o.get("source") == SOURCE:
        return True
    sources = o.get("sources") or []
    if isinstance(sources, list) and SOURCE in sources and o.get("source") != "anterior-segment-overview":
        return True
    return False


def relink_to_feasibility(dry_run: bool = True):
    """Retarget dry-eye outcomes onto feasibility (or best) sites; remove orphan stubs.

    Never overwrites site/study docs beyond optional Dry Eye tags on the *target*.
    Never changes anterior funnel numbers on non-dry-eye outcomes.
    """
    load_key_from_local_settings()
    if not KEY:
        raise SystemExit("COSMOS_KEY missing")
    client = CosmosClient(ENDPOINT, credential=KEY)
    db = client.get_database_client(DATABASE_ID)
    sites_c = db.get_container_client(SITES_CONTAINER)
    outcomes_c = db.get_container_client(OUTCOMES_CONTAINER)

    sites = list(
        sites_c.query_items("SELECT * FROM c", enable_cross_partition_query=True)
    )
    outcomes = list(
        outcomes_c.query_items("SELECT * FROM c", enable_cross_partition_query=True)
    )
    site_by_id = {s["id"]: s for s in sites}
    pi_index = build_pi_site_index(outcomes, sites)

    def resolve_label(label: str) -> dict | None:
        if not label:
            return None
        return match_site(label, [s for s in sites if (s.get("source") or "") != SOURCE], pi_index)

    dry_outs = [o for o in outcomes if is_dry_eye_outcome(o)]
    print(f"Dry-eye-linked outcomes: {len(dry_outs)}")
    print(f"PI index size: {len(pi_index)}")

    retargets = []
    loc_qual_re = re.compile(
        r"\b(MA|RI|NY|CA|TX|FL|PA|NC|Memphis|Andover|Raynham|Loden|cx)\b",
        re.I,
    )
    for o in dry_outs:
        label = o.get("workbookSiteLabel") or o.get("pi") or o.get("siteName")
        hit = resolve_label(label) if label else None
        if not hit:
            continue
        if hit["id"] == o.get("siteId"):
            continue
        # Only move toward feasibility / non-stub; never onto another dry-eye stub
        target = site_by_id.get(hit["id"])
        if not target or (target.get("source") or "") == SOURCE:
            continue
        # Location-qualified labels must land on a site that shares the qualifier
        quals = loc_qual_re.findall(label or "")
        if quals:
            target_blob = " ".join(
                [
                    target.get("name") or "",
                    target.get("city") or "",
                    target.get("state") or "",
                    target.get("pi") or "",
                ]
            ).lower()
            if not any(q.lower() in target_blob for q in quals):
                continue
        cur = site_by_id.get(o.get("siteId") or "")
        cur_w = site_source_weight(cur)
        new_w = site_source_weight(target)
        # Move if target is better for surveys, or current is a dry-eye stub
        if (cur and (cur.get("source") or "") == SOURCE) or new_w > cur_w:
            retargets.append((o, hit))

    print(f"Outcomes to retarget onto better/feasibility sites: {len(retargets)}")
    for o, hit in retargets[:40]:
        print(
            f"  {o.get('workbookSiteLabel') or o.get('pi')!r}: "
            f"{o.get('siteName')!r} -> {hit.get('name')!r} ({hit['id']})"
        )
    if len(retargets) > 40:
        print(f"  ... +{len(retargets) - 40} more")

    # Stub sites that can merge entirely (plain PI surnames only — keep MA/RI/Loden variants)
    loc_qual_re = re.compile(
        r"\b(MA|RI|NY|CA|TX|FL|PA|NC|Memphis|Andover|Raynham|Loden|cx)\b",
        re.I,
    )
    stubs = [s for s in sites if (s.get("source") or "") == SOURCE]
    stub_merges = []
    for stub in stubs:
        sname = stub.get("name") or ""
        hit = resolve_label(sname)
        if not hit or hit["id"] == stub["id"]:
            continue
        quals = loc_qual_re.findall(sname)
        if quals:
            target_blob = " ".join(
                [
                    hit.get("name") or "",
                    (site_by_id.get(hit["id"]) or {}).get("city") or "",
                    (site_by_id.get(hit["id"]) or {}).get("pi") or "",
                ]
            ).lower()
            if not any(q.lower() in target_blob for q in quals):
                continue  # e.g. Jordan (MA) must not merge into Raynham
        stub_merges.append((stub, hit))

    print(f"Dry-eye stub sites mergeable into existing: {len(stub_merges)}")
    for stub, hit in stub_merges:
        print(f"  stub {stub.get('name')!r} -> {hit.get('name')!r}")

    if dry_run:
        print("\nDRY RUN relink — no writes. Re-run with --relink --apply")
        return

    now = datetime.utcnow().isoformat() + "Z"
    touched_targets = set()

    for o, hit in retargets:
        o["siteId"] = hit["id"]
        o["siteName"] = hit.get("name") or o.get("siteName")
        o["updatedAt"] = now
        o["relinkedToFeasibilityAt"] = now
        outcomes_c.upsert_item(o)
        touched_targets.add(hit["id"])

    # Also retarget any remaining outcomes still on mergeable stubs
    for stub, hit in stub_merges:
        for o in outcomes:
            if o.get("siteId") != stub["id"]:
                continue
            o["siteId"] = hit["id"]
            o["siteName"] = hit.get("name") or o.get("siteName")
            o["updatedAt"] = now
            o["relinkedToFeasibilityAt"] = now
            outcomes_c.upsert_item(o)
            touched_targets.add(hit["id"])

    # Annotate targets (additive only)
    for tid in touched_targets:
        full = sites_c.read_item(tid, partition_key=tid)
        changed = False
        inds = full.get("indicationsCovered") or []
        if not isinstance(inds, list):
            inds = [inds] if inds else []
        if INDICATION not in inds:
            inds.append(INDICATION)
            full["indicationsCovered"] = inds
            changed = True
        sources = full.get("sources") or []
        if not isinstance(sources, list):
            sources = [sources] if sources else []
        if full.get("source") and full["source"] not in sources:
            sources.append(full["source"])
        if SOURCE not in sources:
            sources.append(SOURCE)
            full["sources"] = sources
            changed = True
        if changed:
            full["updatedAt"] = now
            sites_c.upsert_item(full)

    # Delete stubs that no longer have outcomes
    remaining = list(
        outcomes_c.query_items(
            "SELECT c.siteId FROM c",
            enable_cross_partition_query=True,
        )
    )
    still_used = {r.get("siteId") for r in remaining if r.get("siteId")}
    deleted = 0
    for stub, hit in stub_merges:
        if stub["id"] not in still_used:
            sites_c.delete_item(stub["id"], partition_key=stub["id"])
            deleted += 1
            print(f"  deleted stub {stub['id']} (merged into {hit['id']})")
    print(f"Relink done. Retargeted outcomes={len(retargets)}, stubs deleted={deleted}")


def main():
    flags = set(sys.argv[1:])
    dry = "--apply" not in flags
    if "--dry-run" in flags:
        dry = True
    do_relink = "--relink" in flags
    argv = [a for a in sys.argv[1:] if not a.startswith("--")]

    if do_relink and not argv:
        # Relink-only mode
        relink_to_feasibility(dry_run=dry)
        return

    path = Path(argv[0]) if argv else DEFAULT_XLSX
    if not path.exists():
        raise SystemExit(f"File not found: {path}")

    print(f"Parsing {path} ...")
    records, sheet_stats = parse_workbook(path)
    print(f"Parsed {len(records)} site-study rows from {sum(1 for x in sheet_stats if x['rows'])} sheets")
    for st in sheet_stats:
        flag = "OK" if st["rows"] else "SKIP"
        print(f"  {flag:4} {st['sheet'][:40]:40} study={st['study']!r:40} rows={st['rows']}")

    load_key_from_local_settings()
    if not KEY:
        raise SystemExit("COSMOS_KEY missing")
    client = CosmosClient(ENDPOINT, credential=KEY)
    db = client.get_database_client(DATABASE_ID)
    existing_studies, existing_sites, existing_outcomes = fetch_existing(db)

    studies, sites, outcomes, match_report = build_docs(
        records, existing_studies, existing_sites, existing_outcomes, str(path)
    )
    upsert_all(studies, sites, outcomes, match_report, dry_run=dry)
    if dry:
        print("\nRe-run with --apply to write.")
        if do_relink:
            print("Relink skipped during dry-run ingest; run: python ingest/legacy_dry_eye_overview.py --relink --dry-run")
        return

    if do_relink:
        print("\n=== RELINK to feasibility ===")
        relink_to_feasibility(dry_run=False)


if __name__ == "__main__":
    main()
