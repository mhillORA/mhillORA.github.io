"""Ingest Stealth GA Feasibility.xlsx into ReBUILD survey responses.

Stealth wins over Mike pack / prior answers for overlapping library IDs.
Matches every usable Stealth row to a live Artemis site (email > PI+institution).

Usage:
  python ingest/ingest_stealth_ga_rebuild.py
  python ingest/ingest_stealth_ga_rebuild.py --apply
"""
from __future__ import annotations

import argparse
import json
import re
import time
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ingest_mike_rebuild_pack import (
    SURVEY_ID,
    coerce_to_question_value,
    cosmos_db,
    iso_now,
    sanitize_pack_value,
)

REPO = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\shue1\Downloads\Stealth GA Feasibility.xlsx")
REPORT = REPO / ".firecrawl" / "stealth-ga-rebuild-ingest-report.json"
SOURCE = "stealth-ga-feasibility"
MAP_OUT = REPO / ".firecrawl" / "stealth-ga-column-to-rebuild.json"

# Curated Stealth column -> rebuild libraryQuestionId (Response cols + specials)
# Verified against Stealth GA Feasibility.xlsx Q inference + live rebuild survey.
COL_LIB: dict[str, str] = {
    "AS": "ql-rebuild-013-has-the-investigator-reviewed-the-protocol-synop",  # resolved live
    "BB": "ql-gf-05-practice-setting",
    "BD": "ql-rebuild-018-does-your-site-have-satellite-offices-or-other-locat",
    "BO": "ql-rebuild-023-how-many-dry-amd-clinical-trials-has-your-site-condu",
    # BS is Yes/No in Stealth; rebuild asks how-many → handled via BT in build_answers
    "BU": "ql-rebuild-026-if-you-have-ongoing-dry-amd-ga-trials-would-they-int",
    "CC": "ql-rebuild-072-is-your-staff-gcp-certified",
    "GH": "ql-rebuild-077-do-you-have-a-separate-contract-and-or-budget-office",
    "CD": "ql-rebuild-025-do-you-have-dedicated-staff-to-conduct-this-study",
    "CS": "ql-rebuild-030-has-your-site-staff-equipment-ever-been-certified-by",
    "DE": "mu3330hmyp18qfkmzy",  # FAF Yes/No
    "DG": "mu3330pegxxzitpbtm4",  # FA Yes/No
    "DL": "mu3330woingrwq22ky",  # ETDRS lightbox
    "DN": "ql-rebuild-043-do-you-have-a-dedicated-4-meter-lane-room-area-to-pe",
    "DO": "ql-rebuild-055-does-your-site-maintain-a-regular-calibration-schedu",
    "DP": "ql-gsf_075_are-calibration-records-available-for-monitoring",
    "DV": "ql-rebuild-051-will-cra-have-remote-access-to-the-electronic-source",
    "DW": "ql-rebuild-064-on-site-pharmacy",
    "DX": "ql-rebuild-061-would-your-study-staff-have-unrestricted-access-to-a",
    "EI": "ql-rebuild-063-do-you-need-a-separate-trial-specific-pharmacy-contr",
    "EJ": "ql-rebuild-065-2-8-c-secured-refrigerator-with-continuous-temperatu",
    "EK": "ql-rebuild-066-secured-temperature-controlled-storage-room-for-anci",
    "EL": "ql-rebuild-069-20-c-freezer-with-continuous-temperature-monitoring-",
    "EN": "ql-rebuild-067-are-you-able-to-conduct-on-site-blood-draws",
    "ES": "ql-rebuild-070-is-your-site-able-to-obtain-dry-ice-for-shipping-spe",
    "ET": "ql-gsf_072_does-your-site-have-one-or-more-staff-who-are-ia",
    "EU": "ql-rebuild-073-centrifuge",
    "EV": "ql-rebuild-074-ability-to-collect-process-and-ship-package-and-labe",
    "EY": "ql-rebuild-075-secure-space-to-store-clinical-study-files-and-subje",
    "FA": "ql-gsf_073_does-your-site-have-a-backup-plan-for-power-outa",
    "FC": "ql-rebuild-078-what-percentage-of-these-patients-have-extrafoveal-g",
    "FM": "ql-rebuild-079-do-you-use-complement-inhibitors-in-the-studys-propo",  # resolve live
    "FP": "ql-rebuild-060-is-your-site-willing-to-pre-screen-potential-suitabl",
    "FS": "ql-gf-29-central-irb",
    "FT": "ql-rebuild-063-if-local-irb-ec-how-often-does-the-irb-meet",
    "GG": "ql-rebuild-074-can-irb-ec-submission-and-contract-budget-negotiatio",
    "GM": "ql-rebuild-080-on-average-how-long-will-it-take-to-execute-the-clin",
    "GN": "ql-rebuild-081-do-you-accept-electronic-signatures-of-the-clinical-",
}

# CFP capability checkboxes
CFP_COLS = {
    "CZ": "Standard 3-field",
    "DA": "Ultra Wide-Field",
}
CFP_LIB = "mu332z1gb2ehc1rf3n"
CFP_MODEL_COL = "DB"
CFP_MODEL_LIB = "ql-rebuild-044-manufacturer-model-if-yes"

# Source review checkboxes → rebuild radio options
SOURCE_COLS = {
    "DQ": "Paper source or certified copies",
    "DR": "EMR, CRA can log in remotely",
    "DS": "EMR, CRA can review certified printouts",
}
SOURCE_LIB = "ql-gsf_060_how-will-the-monitor-review-source-data-at-your-"
SOURCE_OTHER_COL = "DT"
SOURCE_OTHER_LIB = "ql-rebuild-058-if-other-please-specify"
ESOURCE_SYSTEM_COL = "DU"
ESOURCE_SYSTEM_LIB = "ql-rebuild-059-if-electronic-source-which-system"

# Patient identify methods
PATIENT_ID_COLS = {
    "FF": "Site database review",
    "FG": "Patient chart review",
    "FH": "Dear Dr. letter",
    "FI": "Past enrollment in similar studies",
    "FJ": "Physician referrals",
}
PATIENT_ID_LIB = "ql-rebuild-085-what-method-s-are-you-planning-to-use-to-identify-pa"
PATIENT_ID_OTHER_COL = "FK"
PATIENT_ID_OTHER_LIB = "ql-rebuild-086-if-other-please-specify"

VOLUME_COLS = {
    "FB": "ql-rebuild-077-how-many-patients-over-55-years-old-with-at-least-on",
    "FD": "ql-rebuild-054-how-many-newly-referred-or-newly-diagnosed-with-dry-",
    "FE": "ql-rebuild-080-based-on-the-numbers-above-and-the-protocol-synopsis",
}

EDC_LIB = "ql-gsf_065_please-select-which-of-the-following-edc-systems"
EDC_OTHER_LIB = "ql-rebuild-105-if-other-please-specify"
EDC_TEXT_COL = "GP"

RC_LIB = "ql-gsf_087_please-indicate-the-central-imaging-reading-cent"
RC_OTHER_LIB = "ql-rebuild-109-if-other-please-specify"
RC_TEXT_COL = "GS"
RC_YES_COL = "GR"

BLOOD_WHERE_COL = "EO"
BLOOD_WHERE_LIB = "ql-rebuild-068-if-no-where-will-patients-go-for-blood-draws"

EDC_ALIASES = [
    ("imedidata", "Medidata Rave"),
    ("medidata", "Medidata Rave"),
    ("rave", "Medidata Rave"),
    ("inform", "Inform"),
    ("oracle", "Oracle Clinical"),
    ("redcap", "REDCap"),
    ("imednet", "iMednet"),
    ("openclinica", "OpenClinica"),
]
RC_ALIASES = [
    ("duke", "Duke"),
    ("merit", "MERIT"),
    ("clario", "Clario"),
    ("ert", "Clario"),
    ("oirrc", "OIRRC"),
    ("adaptive sensory", "Adaptive Sensory Technology (AST)"),
    (r"\bast\b", "Adaptive Sensory Technology (AST)"),
    ("birc", "BIRC"),
    ("ciarc", "CIARC"),
]

# When Yes carries model text, model lives in the next column
MODEL_FOLLOWUP: dict[str, tuple[str, str]] = {
    # stealth_col: (model_col, rebuild_model_lib)
    "DE": ("DF", "ql-rebuild-046-manufacturer-model-if-yes"),
    "DG": ("DH", "ql-rebuild-048-manufacturer-model-if-yes"),
    "DL": ("DM", "ql-rebuild-050-make-model-catalog-no-if-yes"),
}

# Checkbox role columns (truthy if cell equals header / Yes / checked)
ROLE_COLS = {
    "CE": "Study Coordinator",
    "CF": "Pharmacist",
    "CG": "Sub-Investigators",
    "CH": "Certified Photographers/Technicians",
    "CI": "Certified BCVA Examiners",
}
ROLES_LIB = "ql-rebuild-034-if-yes-which-dedicated-roles-are-in-place"

# SD-OCT brand checkboxes
OCT_COLS = {
    "CU": "Heidelberg Spectralis",
    "CV": "Zeiss Cirrus",
}
OCT_LIB = "ql-rebuild-041-do-you-have-the-capability-to-perform-spectral-domai"
OCT_OTHER_COL = "CW"  # may be Other specify
OCT_OTHER_LIB = "ql-rebuild-042-if-other-please-specify"

# Board-certified PI / Sub-I checkboxes
BOARD_COLS = {"CK": "PI", "CL": "Sub-I"}
BOARD_LIB = "ql-rebuild-036-are-you-or-any-of-your-sub-investigators-board-certi"

# Count fields
COUNT_COLS = {
    "CJ": "ql-rebuild-027-how-many-sub-investigators-do-you-plan-to-have-invol",
    "CT": "ql-rebuild-032-how-many-certified-bcva-examiners-do-you-plan-to-hav",
}

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_EXCEL_DATE_RE = re.compile(
    r"^\s*\d{4}-\d{1,2}-\d{1,2}(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?\s*$"
)
_QUESTION_VAL_RE = re.compile(
    r"^(does |do |is |are |have |how |what |which |please |if |has |will |would |response$)",
    re.I,
)


def col_idx(letter: str) -> int:
    return column_index_from_string(letter) - 1


def norm_name(s: str) -> str:
    t = re.sub(r"\s+", " ", str(s or "").lower())
    t = re.sub(
        r"\b(inc\.?|llc|corp\.?|ltd\.?|pllc|pc|md|do|phd|od|dr\.?|sc|the|center|institute|associates|assoc)\b",
        "",
        t,
    )
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def norm_email(s: str) -> str:
    return str(s or "").strip().lower()


def is_junk_value(v) -> bool:
    if v is None:
        return True
    s = str(v).strip()
    if not s or s.lower() in ("response", "open-ended response", "n/a", "na", "null", "none", "-"):
        return True
    if _EXCEL_DATE_RE.match(s) or ("00:00:00" in s and re.search(r"\d{4}-\d{2}-\d{2}", s)):
        return True
    if _QUESTION_VAL_RE.match(s) and ("?" in s or len(s) > 40):
        return True
    return False


def checkbox_checked(cell_val, option_label: str) -> bool:
    if cell_val is None:
        return False
    s = str(cell_val).strip()
    if not s or is_junk_value(s):
        return False
    if re.match(r"^(yes|y|true|1|x|✓|✔)$", s, re.I):
        return True
    ol = option_label.lower()
    sl = s.lower()
    if sl == ol or ol in sl or sl in ol:
        return True
    return False


def yes_no_from_stealth(val) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip()
    if re.match(r"^(yes)\b", s, re.I):
        return "Yes"
    if re.match(r"^(no)\b", s, re.I):
        return "No"
    if re.match(r"^(n/?a|not applicable)$", s, re.I):
        return "N/A"
    # "Yes- please list..." / "Yes- please specify..."
    if re.match(r"^yes\b", s, re.I):
        return "Yes"
    return None


def centrifuge_coerce(val) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip().lower()
    if s in ("yes", "standard", "refrigerated", "both") or "centrifuge" in s:
        return "Yes"
    if s == "no":
        return "No"
    yn = yes_no_from_stealth(val)
    return yn


def practice_coerce(val, opts: list) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip()
    # normalize curly apostrophe
    s2 = s.replace("\u2019", "'").replace("\u2018", "'")
    for o in opts:
        if s2.lower() == o.lower():
            return o
    aliases = {
        "doctor's office (group practice)": "Doctor's Office (Group Practice)",
        "doctor's office (private practice)": "Doctor's Office (Private Practice)",
        "university hospital": "University Hospital",
        "general hospital": "General Hospital",
        "dedicated research center": "Dedicated Research Center",
        "specialized clinic/institution": "Specialized Clinic/Institution",
        "specialty clinic": "Specialized Clinic/Institution",
    }
    hit = aliases.get(s2.lower())
    if hit and hit in opts:
        return hit
    for o in opts:
        if o.lower() in s2.lower() or s2.lower() in o.lower():
            return o
    return None


def irb_coerce(val, opts: list) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip().lower()
    if "central" in s:
        return next((o for o in opts if "central" in o.lower()), None)
    if "local" in s:
        return next((o for o in opts if "local" in o.lower()), None)
    return None


def load_sheet(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Stealth GA"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows[0], rows[1:]


def resolve_interest_lib(survey_qs: list) -> None:
    """Fix AS interest lib id from live survey labels."""
    for q in survey_qs:
        lab = (q.get("label") or "").lower()
        if "reviewed the protocol synopsis" in lab and q.get("libraryQuestionId"):
            COL_LIB["AS"] = q["libraryQuestionId"]
            return


def resolve_ongoing_libs(survey_qs: list) -> None:
    for q in survey_qs:
        lab = (q.get("label") or "").lower()
        lib = q.get("libraryQuestionId")
        if not lib:
            continue
        if "interfere with recruit" in lab:
            COL_LIB["BU"] = lib
        if "dry amd clinical trials has your site conducted" in lab:
            COL_LIB["BO"] = lib
        # Prefer the "in the study's proposed patient population" complement Q
        if "complement inhibitor" in lab and "proposed patient" in lab:
            COL_LIB["FM"] = lib
        elif "complement inhibitor" in lab and "FM" not in COL_LIB:
            COL_LIB["FM"] = lib
        if "satellite offices" in lab:
            COL_LIB["BD"] = lib
        if "calibration records" in lab and "monitoring" in lab:
            COL_LIB["DP"] = lib
        if "separate contract" in lab and "budget" in lab:
            COL_LIB["GH"] = lib


ONGOING_LIB = "ql-rebuild-025-how-many-ongoing-trials-does-your-site-have-in-patie"
PHASE_LIB = "ql-rebuild-024-which-development-phase-s-does-that-experience-cover"
PHASE_COLS = {"BP": "Phase 1", "BQ": "Phase 2", "BR": "Phase 3"}
STUDY_ENROLL = {
    "BV": "ql-rebuild-027-study-1-number-of-subjects-enrolled",
    "BX": "ql-rebuild-029-study-2-number-of-subjects-enrolled",
    "BZ": "ql-rebuild-031-study-3-number-of-subjects-enrolled",
}
STUDY_LENGTH = {
    "BW": "ql-rebuild-028-study-1-length-of-enrollment-period",
    "BY": "ql-rebuild-030-study-2-length-of-enrollment-period",
    "CA": "ql-rebuild-032-study-3-length-of-enrollment-period",
}
COMPLEMENT_PCT_LIB = "ql-rebuild-059-if-yes-what-percentage-of-dry-amd-patients"
COMPLEMENT_ROUTINE_LIB = "ql-rebuild-057-does-your-site-routinely-use-complement-inhibitors-t"
DIARY_EXP_LIB = "ql-rebuild-087-do-you-have-experience-using-paper-patient-diaries-i"
DIARY_PREF_LIB = "ql-rebuild-111-what-type-of-diary-collection-are-your-patients-most"
DIARY_COLS = {"GW": "eDiary", "GX": "Paper Diary"}
IRT_LIB = "ql-rebuild-106-please-select-which-of-the-following-irt-rtsm-system"
IRT_OTHER_LIB = "ql-rebuild-107-if-other-please-specify"
SAT_NAME_LIB = "ql-rebuild-020-satellite-institution-name"


def bucket_ongoing_count(n: int, opts: list[str]) -> str | None:
    if n <= 0:
        target = "0"
    elif n <= 2:
        target = "1-2"
    elif n <= 5:
        target = "3-5"
    else:
        target = ">=6"
    # Match live option punctuation (en-dash / ≥)
    for o in opts:
        norm = o.replace("–", "-").replace("—", "-").replace("≥", ">=").replace(" ", "")
        if norm == target.replace(" ", "") or (target == ">=6" and ("≥6" in o or ">=6" in o or o.strip() == "6+")):
            return o
        if target == "0" and o.strip() == "0":
            return o
    return None


def bucket_enrollment_months(val, opts: list[str]) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip()
    if re.match(r"^(o|n/?a|na)$", s, re.I):
        return None
    m = re.match(r"^(\d+(?:\.\d+)?)\s*(?:months?)?$", s, re.I)
    if not m:
        # already an option-ish label
        for o in opts:
            if o.lower() == s.lower():
                return o
        return None
    months = float(m.group(1))
    if months < 6:
        want = "< 6"
    elif months <= 12:
        want = "6"
    elif months <= 18:
        want = "13"
    else:
        want = "> 18"
    for o in opts:
        ol = o.replace("–", "-").replace("—", "-")
        if want == "< 6" and ol.startswith("<"):
            return o
        if want == "6" and ("6-12" in ol.replace(" ", "") or "6–12" in o):
            return o
        if want == "13" and ("13-18" in ol.replace(" ", "") or "13–18" in o):
            return o
        if want == "> 18" and ol.startswith(">"):
            return o
    return None


def format_address(raw) -> str | None:
    street = str(raw[col_idx("G")] or "").strip() if col_idx("G") < len(raw) else ""
    street2 = str(raw[col_idx("H")] or "").strip() if col_idx("H") < len(raw) else ""
    city = str(raw[col_idx("I")] or "").strip() if col_idx("I") < len(raw) else ""
    state = str(raw[col_idx("J")] or "").strip() if col_idx("J") < len(raw) else ""
    zipc = str(raw[col_idx("K")] or "").strip() if col_idx("K") < len(raw) else ""
    junk = {"", "-", "n/a", "na", "address", "address 2", "city/town", "state/province", "zip/postal code", "x", "xx", "xxx"}
    parts = []
    for p in (street, street2):
        if p.lower() not in junk and not re.fullmatch(r"[xX.\-]+", p):
            parts.append(p)
    city_state = ", ".join(
        [p for p in (city, state) if p.lower() not in junk]
    )
    if city_state:
        parts.append(city_state)
    if zipc.lower() not in junk:
        if parts:
            parts[-1] = f"{parts[-1]} {zipc}".strip() if city_state else zipc
        else:
            parts.append(zipc)
    out = ", ".join(parts).strip(" ,")
    return out or None


def address_parts(raw) -> dict:
    junk = {
        "",
        "-",
        "n/a",
        "na",
        "address",
        "address 2",
        "city/town",
        "state/province",
        "zip/postal code",
        "x",
        "xx",
        "xxx",
    }

    def clean(letter):
        if col_idx(letter) >= len(raw):
            return ""
        s = str(raw[col_idx(letter)] or "").strip()
        if s.lower() in junk or re.fullmatch(r"[xX.\-]+", s):
            return ""
        return s

    a1 = clean("G")
    a2 = clean("H")
    # House number alone on G + street name on H → one street line (keep suite/unit on H)
    if re.fullmatch(r"\d{1,6}[A-Za-z]?", a1 or "") and a2:
        if not re.match(
            r"^(suite|ste\.?|apt\.?|apartment|unit|#|bldg\.?|building|floor|fl\.?)\b",
            a2,
            re.I,
        ):
            a1 = f"{a1} {a2}".strip()
            a2 = ""

    return {
        "address1": a1,
        "address2": a2,
        "city": clean("I"),
        "state": clean("J"),
        "zip": clean("K"),
    }


def parse_edc_systems(text: str, opts: list[str]) -> tuple[list[str], str | None]:
    low = text.lower()
    matched = []
    for needle, opt in EDC_ALIASES:
        if needle in low and opt in opts and opt not in matched:
            matched.append(opt)
    leftover = text.strip()
    if matched and not leftover:
        return matched, None
    # If we matched known systems, still keep leftover as Other specify when extra junk
    if matched:
        return matched, leftover if len(leftover) > 3 and not all(
            n in low for n, _ in EDC_ALIASES if n in low
        ) else None
    return (["Other"] if "Other" in opts else []), leftover


def parse_reading_centers(text: str, opts: list[str]) -> tuple[list[str], str | None]:
    low = text.lower()
    matched = []
    for needle, opt in RC_ALIASES:
        if needle.startswith("\\b"):
            if re.search(needle, low) and opt in opts and opt not in matched:
                matched.append(opt)
        elif needle in low and opt in opts and opt not in matched:
            matched.append(opt)
    if matched:
        return matched, text.strip() if len(text.strip()) > 40 else None
    if re.search(r"none|n/?a", low):
        none = next((o for o in opts if "none" in o.lower()), None)
        return ([none] if none else []), None
    return [], text.strip() or None


def volume_clean(val) -> str | None:
    if is_junk_value(val):
        return None
    s = str(val).strip()
    if _EXCEL_DATE_RE.match(s) or ("00:00:00" in s and re.search(r"\d{4}-\d{2}-\d{2}", s)):
        return None
    if re.match(r"^(xxx+|n/?a|na|o)$", s, re.I):
        return None
    return s


def parse_stealth_rows(hdr, data):
    rows = []
    for r in data:
        pi = str(r[col_idx("E")] or "").strip()
        inst = str(r[col_idx("F")] or "").strip()
        email = str(r[col_idx("M")] or "").strip()
        if not pi and not inst:
            continue
        if pi.lower() in ("investigator name", "at", "name"):
            continue
        # skip embedded header rows (GCP question text in CC)
        cc = str(r[col_idx("CC")] or "")
        if "gcp certified" in cc.lower() or cc.strip() == "Response":
            continue
        es = str(r[col_idx("ES")] or "")
        if "dry ice" in es.lower():
            continue
        rows.append(
            {
                "pi": pi,
                "inst": inst,
                "email": email,
                "phone": str(r[col_idx("N")] or "").strip(),
                "city": str(r[col_idx("I")] or "").strip(),
                "state": str(r[col_idx("J")] or "").strip(),
                "address": str(r[col_idx("G")] or "").strip(),
                "zip": str(r[col_idx("K")] or "").strip(),
                "raw": r,
            }
        )
    return rows


def _pi_last_first(pi: str) -> tuple[str, str]:
    t = re.sub(r"\b(md|do|phd|od|dr\.?|,)\b", " ", str(pi or ""), flags=re.I)
    parts = [p for p in re.split(r"\s+", t.strip()) if p]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0].lower(), ""
    return parts[-1].lower(), parts[0].lower()


def match_site(row, live_sites, by_email: dict):
    email = norm_email(row["email"])
    if email and _EMAIL_RE.match(email) and email in by_email:
        return by_email[email], "email", 1.0

    inst_n = norm_name(row["inst"])
    pi_n = norm_name(row["pi"])
    pi_last, pi_first = _pi_last_first(row["pi"])
    best = None
    for site in live_sites:
        sn = norm_name(site.get("name") or "")
        sp = norm_name(site.get("piName") or site.get("pi") or "")
        sp_last, sp_first = _pi_last_first(site.get("piName") or site.get("pi") or "")
        score = 0.0
        how = ""
        if inst_n and sn:
            if inst_n == sn:
                score, how = 1.0, "exact-name"
            elif len(inst_n) >= 8 and len(sn) >= 8 and (inst_n in sn or sn in inst_n):
                shorter, longer = (inst_n, sn) if len(inst_n) <= len(sn) else (sn, inst_n)
                if len(shorter) / max(len(longer), 1) >= 0.5:
                    score, how = 0.93, "contains-name"
            else:
                r = SequenceMatcher(None, inst_n, sn).ratio()
                if r >= 0.88:
                    score, how = r, "fuzzy-name"
        pi_hit = False
        if pi_n and sp and pi_n == sp:
            pi_hit = True
        elif pi_last and sp_last and pi_last == sp_last and (
            not pi_first or not sp_first or pi_first[0] == sp_first[0]
        ):
            pi_hit = True
        if pi_hit:
            if score >= 0.72:
                score, how = max(score, 0.97), (how or "name") + "+pi"
            else:
                tokens = [t for t in inst_n.split() if len(t) >= 4]
                if any(t in sn for t in tokens):
                    score, how = 0.90, "pi+token"
                elif score >= 0.80:
                    score, how = 0.90, how + "+pi"
        # Require either strong name or name+PI — never PI-only across orgs
        if score >= 0.88 and (best is None or score > best[0]):
            best = (score, site, how)
    if best:
        return best[1], best[2], round(best[0], 3)
    return None, None, 0.0


_PHONE_JUNK_RE = re.compile(r"^(n/?a|na|none|null|unknown|-)$", re.I)


def clean_phone(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s or _PHONE_JUNK_RE.match(s) or is_junk_value(s):
        return None
    s = re.sub(r"_x000[dD]_", "", s).strip()
    digits = re.sub(r"\D", "", s)
    if len(digits) < 7:
        return None
    # Normalize US NANP to 555-555-5555 (keep extension as " x ####")
    ext = ""
    ext_m = re.search(r"(?:ext\.?|extension|x)\s*[:.]?\s*(\d{1,8})\s*$", s, re.I)
    if ext_m:
        ext = ext_m.group(1)
        s = s[: ext_m.start()].strip()
        digits = re.sub(r"\D", "", s)
    d = digits
    if len(d) == 11 and d.startswith("1"):
        d = d[1:]
    if len(d) == 10:
        formatted = f"{d[0:3]}-{d[3:6]}-{d[6:10]}"
        return f"{formatted} x {ext}" if ext else formatted
    return s if not ext else f"{s} x {ext}"


def clean_email(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s or _PHONE_JUNK_RE.match(s) or is_junk_value(s):
        return None
    if not _EMAIL_RE.match(s):
        return None
    return s


def clean_person_name(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s or is_junk_value(s):
        return None
    if s.lower() in ("name", "company", "institution name"):
        return None
    if re.match(r"^(yes|no|response)$", s, re.I):
        return None
    return s


def extract_contacts(raw) -> dict:
    """PI / coordinator / contracts / pharmacy from Stealth contact blocks."""
    return {
        "piName": clean_person_name(raw[col_idx("E")] if col_idx("E") < len(raw) else None),
        "piEmail": clean_email(raw[col_idx("M")] if col_idx("M") < len(raw) else None),
        "piPhone": clean_phone(raw[col_idx("N")] if col_idx("N") < len(raw) else None),
        # Block after PI = Primary Research POC / coordinator
        "coordName": clean_person_name(raw[col_idx("O")] if col_idx("O") < len(raw) else None),
        "coordEmail": clean_email(raw[col_idx("W")] if col_idx("W") < len(raw) else None),
        "coordPhone": clean_phone(raw[col_idx("X")] if col_idx("X") < len(raw) else None),
        # Second contact block = contracting/budgeting
        "contractsName": clean_person_name(raw[col_idx("Y")] if col_idx("Y") < len(raw) else None),
        "contractsEmail": clean_email(raw[col_idx("AG")] if col_idx("AG") < len(raw) else None),
        "contractsPhone": clean_phone(raw[col_idx("AH")] if col_idx("AH") < len(raw) else None),
        # Fallback "main contact for questions" block → fill gaps
        "altName": clean_person_name(raw[col_idx("AI")] if col_idx("AI") < len(raw) else None),
        "altEmail": clean_email(raw[col_idx("AQ")] if col_idx("AQ") < len(raw) else None),
        "altPhone": clean_phone(raw[col_idx("AR")] if col_idx("AR") < len(raw) else None),
        "pharmacyPhone": clean_phone(raw[col_idx("EH")] if col_idx("EH") < len(raw) else None),
    }


# Stealth response columns that have real data but NO matching ReBUILD template question.
# Stored on the response as stealthUnmapped so nothing is dropped for sponsor reporting.
UNMAPPED_STEALTH_COLS = {
    "CQ": "physical_exams_pi_subi_onsite",
    "CX": "octa_capability",
    "CY": "octa_manufacturer_model",
    "DC": "microperimetry_capability",
    "DD": "microperimetry_manufacturer_model",
    "DI": "ecc_machine",
    "DJ": "ecc_manufacturer_model",
    "EM": "freezer_minus70",
    "EP": "patient_stay_up_to_2_hours",
    "EQ": "pk_process_within_30_min",
    "EW": "ecg_onsite_usable",
    "EX": "ecg_alternate_location",
    "EZ": "internet_access",
    "FQ": "genetic_blood_sample_willing",
    "FR": "site_in_us_canada_argentina",
    "GT": "central_ecg_reading_center_experience",
}


def collect_unmapped_stealth(raw) -> dict:
    out = {}
    for letter, key in UNMAPPED_STEALTH_COLS.items():
        idx = col_idx(letter)
        if idx >= len(raw) or is_junk_value(raw[idx]):
            continue
        s = str(raw[idx]).strip()
        # skip question-text residue
        if _QUESTION_VAL_RE.match(s) and ("?" in s or len(s) > 50):
            continue
        out[key] = {"stealthColumn": letter, "value": s}
    return out


def build_answers(row, survey_qs: list) -> list:
    # Always resolve against the live template libs (never mutate the template itself)
    resolve_interest_lib(survey_qs)
    resolve_ongoing_libs(survey_qs)
    q_by_lib = {q.get("libraryQuestionId"): q for q in survey_qs if q.get("libraryQuestionId")}
    raw = row["raw"]
    answers = []
    seen = set()

    def add(lib, value, stealth_col, label_fallback=""):
        if not lib or lib in seen:
            return
        if value is None or (isinstance(value, str) and not str(value).strip()):
            return
        mq = q_by_lib.get(lib) or {}
        label = mq.get("label") or label_fallback or lib
        qtype = mq.get("type") or "text"
        opts = mq.get("options") or []
        clean = sanitize_pack_value(f"stealth.{stealth_col}", value, label, qtype, opts)
        if clean is None:
            return
        coerced = coerce_to_question_value(clean, mq, f"stealth.{stealth_col}")
        if coerced is None or (isinstance(coerced, str) and not str(coerced).strip()):
            return
        seen.add(lib)
        answers.append(
            {
                "questionId": mq.get("id") or lib,
                "libraryQuestionId": lib,
                "label": label,
                "type": qtype,
                "value": coerced if not isinstance(coerced, (dict, list)) else json.dumps(coerced),
                "stealthColumn": stealth_col,
                "source": SOURCE,
            }
        )

    # Identity / phones — always write when Stealth has them
    contacts = extract_contacts(raw)
    add("ql-pi-name", contacts["piName"], "E")
    add("ql-pi-email", contacts["piEmail"], "M")
    add("ql-pi-phone", contacts["piPhone"], "N")
    add("ql-coord-name", contacts["coordName"] or contacts["altName"], "O")
    add("ql-coord-email", contacts["coordEmail"] or contacts["altEmail"], "W")
    add("ql-coord-phone", contacts["coordPhone"] or contacts["altPhone"], "X")
    add("ql-contracts-name", contacts["contractsName"] or contacts["altName"], "Y")
    add("ql-contracts-email", contacts["contractsEmail"] or contacts["altEmail"], "AG")
    # Live ReBUILD template uses gsf_092 for contracts phone (not ql-contracts-phone)
    add(
        "ql-gsf_092_contracting-budgeting-contact-phone-number",
        contacts["contractsPhone"] or contacts["altPhone"],
        "AH",
    )
    # NOTE: do not write ql-contracts-phone — not on live ReBUILD template

    # Institution name (col F)
    inst = row.get("inst") or ""
    if not inst and col_idx("F") < len(raw) and not is_junk_value(raw[col_idx("F")]):
        inst = str(raw[col_idx("F")]).strip()
    if inst and inst.lower() not in ("institution name", "name", "company"):
        add("ql-site-name", inst, "F")

    # Simple mapped response columns
    for letter, lib in COL_LIB.items():
        idx = col_idx(letter)
        if idx >= len(raw):
            continue
        val = raw[idx]
        if is_junk_value(val):
            continue
        mq = q_by_lib.get(lib) or {}
        opts = []
        for o in mq.get("options") or []:
            opts.append(o if isinstance(o, str) else str(o.get("label") or o.get("value") or ""))
        opts = [o for o in opts if o]
        typ = str(mq.get("type") or "").lower()

        if letter == "BB":
            v = practice_coerce(val, opts)
        elif letter == "FS":
            v = irb_coerce(val, opts)
        elif letter == "EU":
            v = centrifuge_coerce(val)
        elif letter == "BO":
            # map 1-2 style / bare counts onto radio options
            s = str(val).strip().replace("–", "-").replace("—", "-")
            if isinstance(val, float) and val == int(val):
                s = str(int(val))
            if is_junk_value(s) or _EXCEL_DATE_RE.match(s) or "00:00:00" in s:
                v = None
            elif s in opts:
                v = s
            elif re.match(r"^\d+(\.0+)?$", s):
                v = bucket_ongoing_count(int(float(s)), opts)
            elif s in ("1-2", "1–2"):
                v = bucket_ongoing_count(1, opts)
            elif s in ("3-5", "3–5"):
                v = bucket_ongoing_count(3, opts)
            elif s in ("0", "≥6", ">=6", "6+"):
                v = bucket_ongoing_count(0 if s == "0" else 6, opts)
            else:
                v = coerce_to_question_value(s, mq, f"stealth.{letter}")
                # last chance: bare digit left after coerce
                if v is not None and re.match(r"^\d+$", str(v)) and str(v) not in opts:
                    v = bucket_ongoing_count(int(v), opts)
        elif letter == "FC":
            s = str(val).strip()
            v = None
            if not is_junk_value(s):
                # Normalize Stealth percent bands → live options
                low = s.lower().replace("≤", "<=").replace("≥", ">=")
                for o in opts:
                    ol = o.lower().replace("≤", "<=").replace("≥", ">=")
                    if s == o or low == ol:
                        v = o
                        break
                if v is None:
                    if re.search(r"25\s*%?\s*to\s*<?=?50| >\s*25.*50", low) or ">25 to" in low or ">25% to" in low:
                        v = next((o for o in opts if "25" in o and "50" in o), None)
                    elif re.search(r"over\s*50|>\s*50|≥\s*50|>=\s*50", low):
                        v = next((o for o in opts if "50" in o and ("over" in o.lower() or ">" in o)), None)
                    elif re.search(r"≤\s*25|<=\s*25|under\s*25|<\s*25|0\s*-\s*25", low) or low in ("25%", "≤25%"):
                        v = next((o for o in opts if "25" in o and "50" not in o), None)
                    else:
                        v = coerce_to_question_value(s, mq, f"stealth.{letter}")
        elif letter == "FT":
            s = str(val).strip()
            yn = None
            if re.match(r"^other\b", s, re.I):
                v = next((o for o in opts if o.lower() == "other"), "Other" if "Other" in opts else s)
            elif s in opts:
                v = s
            else:
                # fuzzy to Weekly / Monthly / etc.
                v = next((o for o in opts if o.lower() == s.lower()), None)
                if v is None:
                    v = coerce_to_question_value(s, mq, f"stealth.{letter}")
        else:
            yn = yes_no_from_stealth(val)
            if yn and typ in ("radio", "select") and opts:
                optset = {o.lower() for o in opts}
                if yn.lower() in optset or (yn == "N/A" and "n/a" in optset):
                    v = next(o for o in opts if o.lower() == yn.lower() or (yn == "N/A" and o.lower() == "n/a"))
                else:
                    v = yn
            elif yn and not opts:
                v = yn
            else:
                # free text / other
                s = str(val).strip()
                # strip "Yes- please..." down to Yes for yes/no; model handled separately
                if re.match(r"^yes\b", s, re.I) and typ in ("radio", "select"):
                    v = "Yes"
                else:
                    v = s
        add(lib, v, letter)

        # model follow-ups
        if letter in MODEL_FOLLOWUP:
            mcol, mlib = MODEL_FOLLOWUP[letter]
            midx = col_idx(mcol)
            if midx < len(raw) and not is_junk_value(raw[midx]):
                model = str(raw[midx]).strip()
                if not re.match(r"^(yes|no|response)", model, re.I):
                    add(mlib, model, mcol)
            else:
                # model sometimes embedded after Yes-
                s = str(val).strip()
                m = re.match(r"^yes[\s\-–:]+(?:please\s+)?(?:list|specify)[^\w]*(.*)$", s, re.I)
                if m and m.group(1).strip():
                    add(MODEL_FOLLOWUP[letter][1], m.group(1).strip(), letter + ".embedded")

    # Dedicated roles multiselect
    roles = []
    for letter, opt in ROLE_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and checkbox_checked(raw[idx], opt):
            # Stealth uses "Certified Photographers/OCT Technicians" -> rebuild "Certified Photographers/Technicians"
            if "Photograph" in opt:
                roles.append("Certified Photographers/Technicians")
            else:
                roles.append(opt)
    if roles:
        add(ROLES_LIB, json.dumps(roles) if len(roles) > 1 else roles[0], "CE-CI")
        if "ql-rebuild-025-do-you-have-dedicated-staff-to-conduct-this-study" not in seen:
            add(
                "ql-rebuild-025-do-you-have-dedicated-staff-to-conduct-this-study",
                "Yes",
                "CE-CI.infer",
            )

    # SD-OCT brands
    octs = []
    for letter, opt in OCT_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and checkbox_checked(raw[idx], opt):
            octs.append(opt)
    if len(octs) >= 2:
        add(OCT_LIB, "Both", "CU-CV")
    elif len(octs) == 1:
        add(OCT_LIB, octs[0], "CU-CV")
    # other specify
    oidx = col_idx(OCT_OTHER_COL) if OCT_OTHER_COL else -1
    if oidx >= 0 and oidx < len(raw) and not is_junk_value(raw[oidx]):
        other = str(raw[oidx]).strip()
        if other.lower() not in ("other", "other (please specify)", "us"):
            if OCT_LIB not in seen:
                add(OCT_LIB, "Other", OCT_OTHER_COL)
            add(OCT_OTHER_LIB, other, OCT_OTHER_COL)

    # Board certified
    boards = []
    for letter, opt in BOARD_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and checkbox_checked(raw[idx], opt):
            boards.append(opt)
    if boards:
        add(BOARD_LIB, json.dumps(boards) if len(boards) > 1 else boards[0], "CK-CL")

    # Counts
    for letter, lib in COUNT_COLS.items():
        idx = col_idx(letter)
        if idx >= len(raw) or is_junk_value(raw[idx]):
            continue
        s = str(raw[idx]).strip()
        if re.match(r"^\d+$", s):
            add(lib, s, letter)

    # Ongoing Dry AMD/GA trial COUNT (Stealth BS=Yes/No, BT=how many)
    mq_ong = q_by_lib.get(ONGOING_LIB) or {}
    ong_opts = [
        o if isinstance(o, str) else str(o.get("label") or o.get("value") or "")
        for o in (mq_ong.get("options") or [])
    ]
    ong_opts = [o for o in ong_opts if o]
    bt_idx = col_idx("BT")
    bs_idx = col_idx("BS")
    ongoing_n = None
    if bt_idx < len(raw) and not is_junk_value(raw[bt_idx]):
        s = str(raw[bt_idx]).strip()
        if re.match(r"^\d+$", s):
            ongoing_n = int(s)
    if ongoing_n is None and bs_idx < len(raw):
        yn = yes_no_from_stealth(raw[bs_idx])
        if yn == "No":
            ongoing_n = 0
        # Yes without BT → leave blank rather than invent a count
    if ongoing_n is not None:
        bucketed = bucket_ongoing_count(ongoing_n, ong_opts)
        if bucketed:
            add(ONGOING_LIB, bucketed, "BT" if ongoing_n else "BS")

    # Development phases (checkboxes)
    phases = []
    for letter, opt in PHASE_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and checkbox_checked(raw[idx], opt):
            phases.append(opt)
    if phases:
        add(PHASE_LIB, json.dumps(phases) if len(phases) > 1 else phases[0], "BP-BR")

    # Prior study enrollment #1-3
    for letter, lib in STUDY_ENROLL.items():
        idx = col_idx(letter)
        if idx >= len(raw) or is_junk_value(raw[idx]):
            continue
        s = str(raw[idx]).strip()
        if re.match(r"^(o|n/?a|na|-)$", s, re.I):
            continue
        if re.match(r"^\d+$", s):
            add(lib, s, letter)
    for letter, lib in STUDY_LENGTH.items():
        idx = col_idx(letter)
        if idx >= len(raw):
            continue
        mq = q_by_lib.get(lib) or {}
        opts = [
            o if isinstance(o, str) else str(o.get("label") or o.get("value") or "")
            for o in (mq.get("options") or [])
        ]
        opts = [o for o in opts if o]
        bucketed = bucket_enrollment_months(raw[idx], opts)
        if bucketed:
            add(lib, bucketed, letter)

    # Satellite institution name
    be_idx = col_idx("BE")
    if be_idx < len(raw) and not is_junk_value(raw[be_idx]):
        sat = str(raw[be_idx]).strip()
        if sat.lower() not in (
            "institution name",
            "alth",
            "satellite office / other location where study procedures will be performed",
            "company",
            "na",
            "n/a",
        ):
            add(SAT_NAME_LIB, sat, "BE")

    # Complement % + routine Yes if they use in proposed population
    fo_idx = col_idx("FO")
    if fo_idx < len(raw) and not is_junk_value(raw[fo_idx]):
        pct = str(raw[fo_idx]).strip()
        if pct.lower() not in ("open-ended response",):
            add(COMPLEMENT_PCT_LIB, pct, "FO")
    if COL_LIB.get("FM") and COL_LIB["FM"] in seen:
        # If they answered the proposed-population complement Q Yes, also mark routine use
        fm_ans = next((a for a in answers if a.get("libraryQuestionId") == COL_LIB["FM"]), None)
        if fm_ans and str(fm_ans.get("value") or "").lower() == "yes":
            add(COMPLEMENT_ROUTINE_LIB, "Yes", "FM.infer")

    # Patient diaries (checkbox pair)
    diary_hits = []
    for letter, opt in DIARY_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and (
            checkbox_checked(raw[idx], opt)
            or (not is_junk_value(raw[idx]) and opt.lower() in str(raw[idx]).strip().lower())
        ):
            if opt not in diary_hits:
                diary_hits.append(opt)
    if diary_hits:
        # Experience Q uses "Paper diaries" / "eDiaries"
        exp_map = {"eDiary": "eDiaries", "Paper Diary": "Paper diaries"}
        exp_vals = [exp_map.get(x, x) for x in diary_hits]
        add(DIARY_EXP_LIB, json.dumps(exp_vals) if len(exp_vals) > 1 else exp_vals[0], "GW-GX")
        add(DIARY_PREF_LIB, json.dumps(diary_hits) if len(diary_hits) > 1 else diary_hits[0], "GW-GX.pref")

    # IRT experience (Stealth Yes/No only — no system list)
    gq_idx = col_idx("GQ")
    if gq_idx < len(raw) and not is_junk_value(raw[gq_idx]):
        yn = yes_no_from_stealth(raw[gq_idx])
        mq = q_by_lib.get(IRT_LIB) or {}
        opts = [
            o if isinstance(o, str) else str(o.get("label") or o.get("value") or "")
            for o in (mq.get("options") or [])
        ]
        if yn == "No" and any("none" in o.lower() for o in opts):
            none = next(o for o in opts if "none" in o.lower())
            add(IRT_LIB, none, "GQ")
        elif yn == "Yes" and "Other" in opts:
            add(IRT_LIB, "Other", "GQ")
            add(IRT_OTHER_LIB, "Stealth: IRT experience confirmed; system not specified", "GQ.other")

    # Institution address
    addr = format_address(raw)
    add("ql-site-address", addr, "G-K")

    # CFP capability + model
    cfp_hit = False
    for letter, opt in CFP_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and checkbox_checked(raw[idx], opt):
            cfp_hit = True
    cfp_model = None
    midx = col_idx(CFP_MODEL_COL)
    if midx < len(raw) and not is_junk_value(raw[midx]):
        cand = str(raw[midx]).strip()
        if cand.lower() not in ("other", "other (please specify)"):
            cfp_model = cand
            cfp_hit = True
    if cfp_hit:
        add(CFP_LIB, "Yes", "CZ-DA")
    if cfp_model:
        add(CFP_MODEL_LIB, cfp_model, CFP_MODEL_COL)

    # Source review method (checkboxes)
    source_hits = []
    for letter, opt in SOURCE_COLS.items():
        idx = col_idx(letter)
        if idx >= len(raw) or is_junk_value(raw[idx]):
            continue
        s = str(raw[idx]).strip().lower()
        if letter == "DQ" and "paper" in s:
            source_hits.append(opt)
        elif letter == "DR" and ("log-in" in s or "login" in s or "electronic" in s):
            source_hits.append(opt)
        elif letter == "DS" and "printout" in s:
            source_hits.append(opt)
    other_src = None
    oidx = col_idx(SOURCE_OTHER_COL)
    if oidx < len(raw) and not is_junk_value(raw[oidx]):
        other_src = str(raw[oidx]).strip()
        if other_src.lower() in ("other", "other (please specify)"):
            other_src = None
    if len(source_hits) == 1:
        add(SOURCE_LIB, source_hits[0], "DQ-DS")
    elif len(source_hits) > 1:
        # Prefer EMR remote login when present
        preferred = next((h for h in source_hits if "log in" in h.lower()), source_hits[-1])
        add(SOURCE_LIB, preferred, "DQ-DS")
    elif other_src:
        add(SOURCE_LIB, "Other", SOURCE_OTHER_COL)
    if other_src:
        add(SOURCE_OTHER_LIB, other_src, SOURCE_OTHER_COL)

    # Electronic source system name
    eidx = col_idx(ESOURCE_SYSTEM_COL)
    if eidx < len(raw) and not is_junk_value(raw[eidx]):
        esys = str(raw[eidx]).strip()
        if esys.lower() not in ("0", "o", "n/a", "na") and not re.match(r"^yes\b", esys, re.I):
            add(ESOURCE_SYSTEM_LIB, esys, ESOURCE_SYSTEM_COL)

    # Patient identify methods
    pid_hits = []
    for letter, opt in PATIENT_ID_COLS.items():
        idx = col_idx(letter)
        if idx < len(raw) and (
            checkbox_checked(raw[idx], opt)
            or (
                not is_junk_value(raw[idx])
                and opt.lower() in str(raw[idx]).strip().lower()
            )
        ):
            if opt not in pid_hits:
                pid_hits.append(opt)
    pid_other = None
    poidx = col_idx(PATIENT_ID_OTHER_COL)
    if poidx < len(raw) and not is_junk_value(raw[poidx]):
        pid_other = str(raw[poidx]).strip()
        if pid_other.lower() in ("other", "other (please specify)"):
            pid_other = None
    if pid_other and "Other" not in pid_hits:
        pid_hits.append("Other")
    if pid_hits:
        add(PATIENT_ID_LIB, json.dumps(pid_hits) if len(pid_hits) > 1 else pid_hits[0], "FF-FJ")
    if pid_other:
        add(PATIENT_ID_OTHER_LIB, pid_other, PATIENT_ID_OTHER_COL)

    # Volume
    for letter, lib in VOLUME_COLS.items():
        idx = col_idx(letter)
        if idx >= len(raw):
            continue
        v = volume_clean(raw[idx])
        if v:
            add(lib, v, letter)

    # Blood draw alternate location
    bidx = col_idx(BLOOD_WHERE_COL)
    if bidx < len(raw) and not is_junk_value(raw[bidx]):
        bw = str(raw[bidx]).strip()
        if not bw.lower().startswith("no - where") and len(bw) > 3:
            add(BLOOD_WHERE_LIB, bw, BLOOD_WHERE_COL)

    # Independent pharmacy contact blob
    pharm_name = clean_person_name(raw[col_idx("DY")] if col_idx("DY") < len(raw) else None)
    pharm_email = clean_email(raw[col_idx("EG")] if col_idx("EG") < len(raw) else None)
    pharm_phone = clean_phone(raw[col_idx("EH")] if col_idx("EH") < len(raw) else None)
    pharm_co = None
    if col_idx("DZ") < len(raw) and not is_junk_value(raw[col_idx("DZ")]):
        cand = str(raw[col_idx("DZ")]).strip()
        if cand.lower() not in ("company",):
            pharm_co = cand
    pharm_addr_parts = []
    for letter in ("EA", "EB", "EC", "ED", "EE", "EF"):
        if col_idx(letter) < len(raw) and not is_junk_value(raw[col_idx(letter)]):
            s = str(raw[col_idx(letter)]).strip()
            if s.lower() not in (
                "company",
                "address",
                "address 2",
                "city/town",
                "state/province",
                "zip/postal code",
                "country",
                "na",
                "n/a",
            ):
                pharm_addr_parts.append(s)
    if pharm_name or pharm_email or pharm_phone or pharm_addr_parts:
        # Skip notes that are not real independent-pharmacy contacts
        name_low = (pharm_name or "").lower()
        if re.search(
            r"not\s+(an\s+)?independent|on[- ]site pharmacy set up|^n\.?a\.?\b|^na\b|error\s*-",
            name_low,
        ) and not pharm_email and not pharm_phone:
            pharm_name = None
        if not (pharm_name or pharm_email or pharm_phone or pharm_addr_parts):
            pass
        else:
            blob = "; ".join(
                x
                for x in [
                    f"Name: {pharm_name}" if pharm_name else "",
                    f"Company: {pharm_co}" if pharm_co else "",
                    f"Email: {pharm_email}" if pharm_email else "",
                    f"Phone: {pharm_phone}" if pharm_phone else "",
                    f"Address: {', '.join(pharm_addr_parts)}" if pharm_addr_parts else "",
                ]
                if x
            )
            # Bypass sanitize_pack_value — question label contains Email/Phone/Address and
            # would reject a multi-field contact blob as "not a valid email/phone".
            pharm_lib = "ql-rebuild-062-if-pharmacy-is-independent-to-your-site-please-provi"
            if blob and pharm_lib not in seen:
                mq = q_by_lib.get(pharm_lib) or {}
                seen.add(pharm_lib)
                answers.append(
                    {
                        "questionId": mq.get("id") or pharm_lib,
                        "libraryQuestionId": pharm_lib,
                        "label": mq.get("label") or pharm_lib,
                        "type": mq.get("type") or "text",
                        "value": blob,
                        "stealthColumn": "DY-EH",
                        "source": SOURCE,
                    }
                )

    # EDC systems from free text
    edc_idx = col_idx(EDC_TEXT_COL)
    if edc_idx < len(raw) and not is_junk_value(raw[edc_idx]):
        edc_text = str(raw[edc_idx]).strip()
        mq = q_by_lib.get(EDC_LIB) or {}
        opts = [
            o if isinstance(o, str) else str(o.get("label") or o.get("value") or "")
            for o in (mq.get("options") or [])
        ]
        opts = [o for o in opts if o]
        matched, other = parse_edc_systems(edc_text, opts)
        if matched:
            add(EDC_LIB, json.dumps(matched) if len(matched) > 1 else matched[0], EDC_TEXT_COL)
        if not matched and edc_text:
            if "Other" in opts:
                add(EDC_LIB, "Other", EDC_TEXT_COL)
            add(EDC_OTHER_LIB, edc_text, EDC_TEXT_COL)
        elif other:
            add(EDC_OTHER_LIB, other, EDC_TEXT_COL + ".extra")

    # Reading centers from free text
    rc_txt_idx = col_idx(RC_TEXT_COL)
    if rc_txt_idx < len(raw) and not is_junk_value(raw[rc_txt_idx]):
        rc_text = str(raw[rc_txt_idx]).strip()
        if rc_text not in ("-",):
            mq = q_by_lib.get(RC_LIB) or {}
            opts = [
                o if isinstance(o, str) else str(o.get("label") or o.get("value") or "")
                for o in (mq.get("options") or [])
            ]
            opts = [o for o in opts if o]
            matched, other = parse_reading_centers(rc_text, opts)
            if matched:
                add(RC_LIB, json.dumps(matched) if len(matched) > 1 else matched[0], RC_TEXT_COL)
                if any("AST" in x or "Adaptive" in x for x in matched):
                    add(
                        "ql-rebuild-031-has-your-site-staff-equipment-ever-been-certified-by",
                        "Yes",
                        RC_TEXT_COL + ".ast",
                    )
            if other and not matched:
                add(RC_OTHER_LIB, other, RC_TEXT_COL)
            elif other and matched:
                add(RC_OTHER_LIB, other, RC_TEXT_COL + ".extra")

    if contacts.get("contractsPhone"):
        add(
            "ql-gsf_092_contracting-budgeting-contact-phone-number",
            contacts["contractsPhone"],
            "AH.gsf",
        )

    contacts["_address"] = address_parts(raw)
    contacts["_addressFormatted"] = addr
    row["_contacts"] = contacts
    row["_unmapped"] = collect_unmapped_stealth(raw)
    return answers


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--xlsx", default=str(XLSX))
    args = ap.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        raise SystemExit(f"Missing {path}")

    hdr, data = load_sheet(path)
    stealth_rows = parse_stealth_rows(hdr, data)
    print(f"Stealth usable rows: {len(stealth_rows)}")

    db = cosmos_db()
    survey = db.get_container_client("site-survey-definitions").read_item(SURVEY_ID, SURVEY_ID)
    survey_qs = survey.get("questions") or []
    resolve_interest_lib(survey_qs)
    resolve_ongoing_libs(survey_qs)

    # Persist effective column map
    MAP_OUT.parent.mkdir(exist_ok=True)
    MAP_OUT.write_text(
        json.dumps(
            {
                "colLib": COL_LIB,
                "roles": ROLE_COLS,
                "oct": OCT_COLS,
                "board": BOARD_COLS,
                "models": MODEL_FOLLOWUP,
                "counts": COUNT_COLS,
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    live = list(
        db.get_container_client("sites").query_items(
            "SELECT c.id, c.name, c.pi, c.piName, c.piEmail FROM c",
            enable_cross_partition_query=True,
        )
    )
    by_email = {}
    for s in live:
        em = norm_email(s.get("piEmail") or "")
        if em and _EMAIL_RE.match(em):
            by_email[em] = s

    now = iso_now()
    matched = []
    unmatched = []
    for row in stealth_rows:
        site, how, score = match_site(row, live, by_email)
        answers = build_answers(row, survey_qs)
        contacts = row.get("_contacts") or extract_contacts(row["raw"])
        rec = {
            "pi": row["pi"],
            "inst": row["inst"],
            "email": row["email"],
            "siteId": site.get("id") if site else None,
            "siteName": site.get("name") if site else None,
            "matchHow": how,
            "matchScore": score,
            "answerCount": len(answers),
            "answersWithLib": sum(1 for a in answers if a.get("libraryQuestionId")),
            "contacts": contacts,
            "stealthUnmapped": row.get("_unmapped") or {},
        }
        if site and score >= 0.88:
            matched.append({**rec, "answers": answers})
        else:
            unmatched.append(rec)

    # Collapse multiple Stealth PIs → one Artemis site: merge answers (later overwrites)
    by_site: dict[str, dict] = {}
    for m in matched:
        sid = m["siteId"]
        if sid not in by_site:
            by_site[sid] = dict(m)
            continue
        existing = {a["libraryQuestionId"]: a for a in by_site[sid]["answers"] if a.get("libraryQuestionId")}
        for a in m["answers"]:
            lib = a.get("libraryQuestionId")
            if lib:
                existing[lib] = a  # Stealth later row wins
        merged = list(existing.values())
        by_site[sid]["answers"] = merged
        by_site[sid]["answerCount"] = len(merged)
        by_site[sid]["answersWithLib"] = len(merged)
        by_site[sid]["pi"] = m["pi"] or by_site[sid]["pi"]
        by_site[sid]["matchHow"] = by_site[sid]["matchHow"] + "+" + (m["matchHow"] or "")
        # contacts: fill blanks from later, overwrite phones when present (Stealth wins)
        prev_c = by_site[sid].get("contacts") or {}
        next_c = m.get("contacts") or {}
        merged_c = dict(prev_c)
        for k, v in next_c.items():
            if v:
                merged_c[k] = v
        by_site[sid]["contacts"] = merged_c
        # unmapped: merge keys, later wins
        prev_u = by_site[sid].get("stealthUnmapped") or {}
        next_u = m.get("stealthUnmapped") or {}
        merged_u = dict(prev_u)
        merged_u.update(next_u)
        by_site[sid]["stealthUnmapped"] = merged_u
    matched_sites = list(by_site.values())

    phone_filled = sum(
        1
        for m in matched_sites
        if (m.get("contacts") or {}).get("piPhone")
        or (m.get("contacts") or {}).get("coordPhone")
        or (m.get("contacts") or {}).get("contractsPhone")
    )
    print(f"Matched rows: {len(matched)} → unique sites: {len(matched_sites)}  Unmatched rows: {len(unmatched)}")
    print(f"Sites with at least one Stealth phone: {phone_filled}/{len(matched_sites)}")
    by_how = {}
    for m in matched:
        by_how[m["matchHow"]] = by_how.get(m["matchHow"], 0) + 1
    print("Match how:", by_how)
    print(
        f"Avg answers/site: {sum(m['answerCount'] for m in matched_sites)/max(len(matched_sites),1):.1f}"
    )

    report = {
        "generatedAt": now,
        "source": SOURCE,
        "surveyId": SURVEY_ID,
        "stealthRows": len(stealth_rows),
        "matchedRows": len(matched),
        "matchedSites": len(matched_sites),
        "unmatched": len(unmatched),
        "sitesWithPhone": phone_filled,
        "matchHow": by_how,
        "unmatchedSample": unmatched[:25],
        "matchedSample": [
            {
                **{k: v for k, v in m.items() if k not in ("answers", "contacts")},
                "phones": {
                    "pi": (m.get("contacts") or {}).get("piPhone"),
                    "coord": (m.get("contacts") or {}).get("coordPhone"),
                    "contracts": (m.get("contacts") or {}).get("contractsPhone"),
                },
            }
            for m in matched_sites[:15]
        ],
    }
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"Wrote {REPORT}")

    if not args.apply:
        print("Dry run only. Re-run with --apply to write Cosmos.")
        return

    rsp_c = db.create_container_if_not_exists(
        id="site-survey-responses",
        partition_key={"paths": ["/siteId"], "kind": "Hash"},
    )
    asg_c = db.create_container_if_not_exists(
        id="site-survey-assignments",
        partition_key={"paths": ["/siteId"], "kind": "Hash"},
    )
    sites_c = db.get_container_client("sites")

    written = 0
    sites_patched = 0
    for m in matched_sites:
        site_id = m["siteId"]
        answers = m["answers"]
        contacts = m.get("contacts") or {}
        rsp_id = f"rsp-stealth-ga-{site_id}"
        asg_id = f"asg-stealth-ga-{site_id}"
        assignment = {
            "id": asg_id,
            "surveyId": SURVEY_ID,
            "siteId": site_id,
            "targetRole": "pi",
            "status": "submitted",
            "source": SOURCE,
            "createdAt": now,
            "updatedAt": now,
            "submittedAt": now,
        }
        response = {
            "id": rsp_id,
            "assignmentId": asg_id,
            "surveyId": SURVEY_ID,
            "siteId": site_id,
            "targetRole": "pi",
            "displayName": m["pi"] or m["inst"],
            "answers": answers,
            "source": SOURCE,
            "stealthInstitution": m["inst"],
            "stealthPi": m["pi"],
            "stealthEmail": m["email"],
            "stealthUnmapped": m.get("stealthUnmapped") or {},
            "matchHow": m["matchHow"],
            "matchScore": m["matchScore"],
            "answerCount": len(answers),
            "createdAt": now,
            "updatedAt": now,
            "submittedAt": now,
        }
        asg_c.upsert_item(assignment)
        rsp_c.upsert_item(response)
        written += 1

        # Site record patch — Stealth wins for phones/emails (sitePrefill beats survey answers)
        try:
            site = sites_c.read_item(site_id, site_id)
        except Exception:
            site = None
        if site:
            patched = False

            def set_field(key, val):
                nonlocal patched
                if not val:
                    return
                if site.get(key) != val:
                    site[key] = val
                    patched = True

            set_field("piPhone", contacts.get("piPhone"))
            set_field("pi_phone", contacts.get("piPhone"))
            if contacts.get("piEmail"):
                set_field("piEmail", contacts.get("piEmail"))
            if contacts.get("piName"):
                # only fill blank PI name — don't clobber a carefully repaired Heier etc. unless blank
                if not str(site.get("piName") or site.get("pi") or "").strip():
                    set_field("piName", contacts["piName"])
                    set_field("pi", contacts["piName"])
            set_field("siteCoordinator", contacts.get("coordName"))
            set_field("siteCoordinatorEmail", contacts.get("coordEmail"))
            set_field("siteCoordinatorPhone", contacts.get("coordPhone"))
            set_field("contractsName", contacts.get("contractsName"))
            set_field("contractsEmail", contacts.get("contractsEmail"))
            set_field("contractsPhone", contacts.get("contractsPhone"))
            # Address — Stealth wins when present (sitePrefill beats survey answers)
            addr = contacts.get("_address") or {}
            if isinstance(addr, dict):
                set_field("address1", addr.get("address1"))
                set_field("address2", addr.get("address2"))
                set_field("city", addr.get("city"))
                set_field("state", addr.get("state"))
                if addr.get("zip"):
                    set_field("zip", addr["zip"])
                    set_field("zipCode", addr["zip"])
                if contacts.get("_addressFormatted"):
                    set_field("address", contacts["_addressFormatted"])
            if patched:
                site["updatedAt"] = now
                site["stealthGaContactSyncedAt"] = now
                sites_c.upsert_item(site)
                sites_patched += 1

    report["status"] = "applied"
    report["written"] = written
    report["sitesPatched"] = sites_patched
    REPORT.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(
        f"Applied: {written} Stealth GA responses; patched {sites_patched} site records with phones/contacts/address."
    )


if __name__ == "__main__":
    main()
