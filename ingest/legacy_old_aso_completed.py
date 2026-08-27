"""
Ingest Completed Studies from Old ASO.xlsx → Cosmos legacy containers.

Separate from Dry Eye Overview. Includes Enrollment Done tabs **and** other
study sheets in the workbook (e.g. YuYu). Skips blank/metrics summary tabs.

PI surnames are the site labels (YuYu: PI column). Match existing legacy
studies/sites 1:1. Outcome ids include this source so they never collide with
anterior / dry-eye rows.

Usage:
  python ingest/legacy_old_aso_completed.py --dry-run
  python ingest/legacy_old_aso_completed.py --apply
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

# Ensure sibling ingest modules import when run as a script
_INGEST_DIR = Path(__file__).resolve().parent
if str(_INGEST_DIR) not in sys.path:
    sys.path.insert(0, str(_INGEST_DIR))
_REPO = _INGEST_DIR.parent
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from azure.cosmos import CosmosClient, PartitionKey
from openpyxl import load_workbook

import legacy_dry_eye_overview as de

DEFAULT_XLSX = Path(r"c:\Users\shue1\Downloads\Completed Studies from Old ASO.xlsx")
SOURCE = "old-aso-completed"
INDICATION = None  # leave blank unless study already has one; don't invent

ENDPOINT = de.ENDPOINT
DATABASE_ID = de.DATABASE_ID
STUDIES_CONTAINER = de.STUDIES_CONTAINER
SITES_CONTAINER = de.SITES_CONTAINER
OUTCOMES_CONTAINER = de.OUTCOMES_CONTAINER

SKIP_SHEETS = {"sheet1", "presbyopia metrics", "allergy metrics"}


def outcome_id(study: str, site: str, group) -> str:
    raw = f"{study}|{site}|{group}|{SOURCE}"
    h = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]
    return f"legacy-outcome-{h}"


def parse_group_cell(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip()
    m = re.search(r"(\d+)", t)
    return float(m.group(1)) if m else None


def map_headers_aso(cells: list) -> dict[str, int] | None:
    """Like dry-eye map_headers, plus explicit Group column."""
    roles = de.map_headers(cells)
    if not roles:
        # Retry softer: Sites + Group + Enrolled without dates (rare)
        raw = [de.s(c) or "" for c in cells]
        norms = [de.norm_name(x) for x in raw]
        soft: dict[str, int] = {}
        for i, n in enumerate(norms):
            if n == "sites":
                soft["site"] = i
            elif n == "group":
                soft["group"] = i
            elif n in {"actual enrolled", "number enrolled", "enrolled"}:
                soft["enrolled"] = i
            elif "schedul" in n and "v1" in n and "target" not in n:
                soft.setdefault("scheduled", i)
            elif "screen" in n and "target" not in n:
                soft.setdefault("screened", i)
            elif "start" in n and "date" in n:
                soft["start"] = i
            elif "end" in n and "date" in n:
                soft["end"] = i
            elif "target" in n and "enroll" in n:
                soft["target_enrolled"] = i
            elif "target" in n and "schedul" in n:
                soft["target_scheduled"] = i
        if "site" in soft and any(k in soft for k in ("scheduled", "screened", "enrolled")):
            roles = soft
        else:
            return None
    # Attach Group col if present
    raw = [de.s(c) or "" for c in cells]
    for i, h in enumerate(raw):
        if de.norm_name(h) == "group":
            roles["group"] = i
            break
    return roles


def resolve_study_label(sheet_name: str, extracted: str | None) -> str:
    """Prefer Study Name cell; keep protocol tokens available for matching."""
    cleaned_tab = re.sub(
        r"\s*-?\s*enrollment\s+done\s*$", "", sheet_name or "", flags=re.I
    ).strip()
    base = (extracted or "").strip() or cleaned_tab
    proto = re.search(
        r"(ADX-\d+|AR\d+|CYS-\d+|Saturn[-\s]?\d+|VELOS-\d+|OPP\s*\d+)",
        sheet_name or "",
        re.I,
    )
    if not proto:
        return base
    token = re.sub(r"\s+", "-", proto.group(1).strip())
    if de.norm_name(token) in de.norm_name(base):
        return base
    # e.g. Study Name "Aerie" + tab "Aerie AR15512-..." → "Aerie (AR15512)"
    # match_study will hit existing Aerie and/or ADX-* via token candidates
    return f"{base} ({token})"


def _looks_like_junk_label(val: str | None) -> bool:
    """Chart/metrics debris from YuYu-style sheets (dates, visit grids, ids)."""
    if not val:
        return True
    t = str(val).strip()
    if not t:
        return True
    low = t.lower()
    if low in {"site", "total", "unique id", "graph date", "month end", "study"}:
        return True
    if re.fullmatch(r"\d+(\.\d+)?", t):
        return True
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
        return True
    if re.match(
        r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\b", low
    ) and re.search(r"\d", t):
        return True
    if "visit" in low and ("day" in low or "cae" in low):
        return True
    return False


def parse_yuyu_flat(rows: list, sheet_name: str) -> list[dict]:
    """YuYu-style flat table: Unique ID | Study | Site | Group | PI | ... | Scheduled | Screen | Randomized."""
    hdr_i = None
    col = {}
    for i, row in enumerate(rows[:40]):
        cells = [de.s(c) for c in (row or [])]
        norms = [de.norm_name(c or "") for c in cells]
        if not norms or norms[0] != "unique id":
            continue
        hdr_i = i
        for j, n in enumerate(norms):
            if n == "study":
                col["study"] = j
            elif n == "site":
                col["site"] = j
            elif n == "group":
                col["group"] = j
            elif n == "pi":
                col["pi"] = j
            elif n in {"fpfv", "visit1 start"}:
                col["start"] = j
            elif n == "lplv":
                col["end"] = j
            elif n == "target scheduled":
                col["target_scheduled"] = j
            elif n == "scheduled":
                col["scheduled"] = j
            elif n in {"screen", "screened"}:
                col["screened"] = j
            elif n in {"randomized", "enrolled", "actual enrolled", "number enrolled"}:
                col["enrolled"] = j
        break
    if hdr_i is None or "pi" not in col or "study" not in col:
        return []

    out = []
    blanks = 0
    for row in rows[hdr_i + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            blanks += 1
            if blanks >= 2 and out:
                break
            continue
        blanks = 0
        uid = de.s(row[0]) if len(row) > 0 else None
        if not uid:
            continue
        if uid.lower() in {"total", "unique id"}:
            break
        # Next block on YuYu is a visit grid whose first col is "Site"
        if uid.lower() == "site" or _looks_like_junk_label(uid):
            break

        study = de.s(row[col["study"]]) if len(row) > col["study"] else None
        pi = de.s(row[col["pi"]]) if len(row) > col["pi"] else None
        clinic = de.s(row[col["site"]]) if "site" in col and len(row) > col["site"] else None
        if _looks_like_junk_label(study) or _looks_like_junk_label(pi):
            break
        if not study or not pi:
            continue
        group = parse_group_cell(row[col["group"]]) if "group" in col and len(row) > col["group"] else None

        def cell(role):
            if role not in col:
                return None
            idx = col[role]
            return row[idx] if len(row) > idx else None

        scheduled = de.num(cell("scheduled"))
        screened = de.num(cell("screened"))
        enrolled = de.num(cell("enrolled"))
        if scheduled is None and screened is None and enrolled is None:
            continue

        # PI is the site identity for matching (same as other Old ASO tabs)
        site_name, g2 = de.parse_group_from_site(pi)
        if group is None:
            group = g2

        out.append(
            {
                "sheet": sheet_name,
                "study": study,
                "site": site_name,
                "site_raw": f"{pi} ({clinic})" if clinic else pi,
                "group": group,
                "pi": site_name,
                "visit1_start": de.s(cell("start")),
                "lplv": de.s(cell("end")),
                "target_scheduled": de.num(cell("target_scheduled")),
                "scheduled": scheduled,
                "screened": screened,
                "enrolled": enrolled,
                "unique_id": uid,
            }
        )
    return out


def prune_junk_old_aso(db, dry_run: bool = True) -> None:
    """Remove chart/metrics debris accidentally ingested as studies/sites/outcomes."""
    _, _, outcomes = de.fetch_existing(db)
    studies_c = db.get_container_client(STUDIES_CONTAINER)
    sites_c = db.get_container_client(SITES_CONTAINER)
    outs_c = db.get_container_client(OUTCOMES_CONTAINER)

    studies = list(
        studies_c.query_items(
            "SELECT * FROM c WHERE c.type = 'legacyStudy'",
            enable_cross_partition_query=True,
        )
    )
    sites = list(
        sites_c.query_items(
            "SELECT * FROM c WHERE c.type = 'legacySite'",
            enable_cross_partition_query=True,
        )
    )

    junk_study_ids = {
        s["id"]
        for s in studies
        if (s.get("source") or "") == SOURCE and _looks_like_junk_label(s.get("name"))
    }
    junk_site_ids = {
        s["id"]
        for s in sites
        if (s.get("source") or "") == SOURCE and _looks_like_junk_label(s.get("name"))
    }
    junk_outs = [
        o
        for o in outcomes
        if (o.get("source") or "") == SOURCE
        and (
            o.get("studyId") in junk_study_ids
            or o.get("siteId") in junk_site_ids
            or _looks_like_junk_label(o.get("studyName"))
            or _looks_like_junk_label(o.get("workbookStudyLabel"))
            or _looks_like_junk_label(o.get("siteName"))
        )
    ]
    print(
        f"Junk Old ASO: studies={len(junk_study_ids)} sites={len(junk_site_ids)} "
        f"outcomes={len(junk_outs)}"
    )
    if dry_run:
        for s in sorted(junk_study_ids):
            print(f"  would delete study {s}")
        for s in sorted(junk_site_ids)[:20]:
            print(f"  would delete site {s}")
        if len(junk_site_ids) > 20:
            print(f"  ... +{len(junk_site_ids) - 20} sites")
        return

    for o in junk_outs:
        outs_c.delete_item(o["id"], partition_key=o["studyId"])
    for sid in junk_study_ids:
        studies_c.delete_item(sid, partition_key=sid)
    for sid in junk_site_ids:
        sites_c.delete_item(sid, partition_key=sid)
    print("Junk prune done.")


def parse_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    records = []
    sheet_stats = []

    for sheet_name in wb.sheetnames:
        low = sheet_name.strip().lower()
        if low in SKIP_SHEETS:
            continue

        rows = list(wb[sheet_name].iter_rows(values_only=True))
        if not rows:
            continue

        # YuYu / flat Unique-ID table
        flat = parse_yuyu_flat(rows, sheet_name)
        if flat:
            records.extend(flat)
            sheet_stats.append(
                {
                    "sheet": sheet_name,
                    "study": flat[0]["study"],
                    "rows": len(flat),
                    "skip": None,
                }
            )
            continue

        study_name = resolve_study_label(sheet_name, de.extract_study_name(rows))

        headers: list[tuple[int, dict]] = []
        for i, row in enumerate(rows[:120]):
            mapped = map_headers_aso(list(row or []))
            if mapped:
                headers.append((i, mapped))
        if not headers:
            sheet_stats.append(
                {"sheet": sheet_name, "study": study_name, "rows": 0, "skip": "no_header"}
            )
            continue

        n = 0
        for hi, (header_i, roles) in enumerate(headers):
            end_at = headers[hi + 1][0] if hi + 1 < len(headers) else min(len(rows), header_i + 80)
            blanks = 0
            for row in rows[header_i + 1 : end_at]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    blanks += 1
                    if blanks >= 3:
                        break
                    continue
                blanks = 0
                if de.is_stop_row(row, roles):
                    site_i = roles["site"]
                    sv = de.s(row[site_i]) if len(row) > site_i else None
                    if sv and sv.lower().startswith("part "):
                        continue
                    break

                site_i = roles["site"]
                site_raw = de.s(row[site_i]) if len(row) > site_i else None
                if not site_raw:
                    continue
                if site_raw.lower() in de.SKIP_SITE_NAMES:
                    continue
                if "average" in site_raw.lower():
                    continue
                if re.fullmatch(r"\d+(\.0+)?", site_raw):
                    continue
                if site_raw.lower().startswith("part "):
                    continue

                site_name, group_from_name = de.parse_group_from_site(site_raw)
                group = group_from_name
                if "group" in roles:
                    g2 = parse_group_cell(row[roles["group"]] if len(row) > roles["group"] else None)
                    if g2 is not None:
                        group = g2

                def cell(role, _roles=roles, _row=row):
                    if role not in _roles:
                        return None
                    idx = _roles[role]
                    return _row[idx] if len(_row) > idx else None

                scheduled = de.num(cell("scheduled"))
                screened = de.num(cell("screened"))
                enrolled = de.num(cell("enrolled"))
                target_scheduled = de.num(cell("target_scheduled"))
                target_enrolled = de.num(cell("target_enrolled"))
                if target_scheduled is None and target_enrolled is not None:
                    target_scheduled = target_enrolled
                start = de.s(cell("start"))
                end = de.s(cell("end"))
                if scheduled is None and screened is None and enrolled is None and start is None:
                    continue

                records.append(
                    {
                        "sheet": sheet_name,
                        "study": study_name,
                        "site": site_name,
                        "site_raw": site_raw,
                        "group": group,
                        "pi": site_name,
                        "visit1_start": start,
                        "lplv": end,
                        "target_scheduled": target_scheduled,
                        "scheduled": scheduled,
                        "screened": screened,
                        "enrolled": enrolled,
                    }
                )
                n += 1

        sheet_stats.append({"sheet": sheet_name, "study": study_name, "rows": n, "skip": None})

    wb.close()
    return records, sheet_stats


def build_docs(records, existing_studies, existing_sites, existing_outcomes, source_file: str):
    now = datetime.utcnow().isoformat() + "Z"
    match_report = {
        "studies_matched": [],
        "studies_new": [],
        "sites_matched": [],
        "sites_new": [],
        "outcomes_new": 0,
        "outcomes_existing_preserved": 0,
    }
    pi_index = de.build_pi_site_index(existing_outcomes, existing_sites)
    matchable_sites = [
        s
        for s in existing_sites
        if (s.get("source") or "") not in {SOURCE, de.SOURCE}
        or (s.get("source") or "").startswith("feasibility")
        or (s.get("source") or "") == "anterior-segment-overview"
    ]
    # Prefer non-stub sites from any prior ingest
    matchable_sites = [
        s
        for s in existing_sites
        if (s.get("source") or "") != SOURCE
    ]

    study_resolve: dict[str, tuple[str, str, bool]] = {}
    site_resolve: dict[str, tuple[str, str, bool]] = {}

    for r in records:
        study = r["study"]
        if study not in study_resolve:
            hit = de.match_study(study, existing_studies)
            if hit:
                study_resolve[study] = (hit["id"], hit.get("name") or study, False)
                match_report["studies_matched"].append(
                    {"from": study, "to": hit.get("name"), "id": hit["id"]}
                )
            else:
                study_resolve[study] = (de.study_id_for(study), study, True)
                match_report["studies_new"].append({"name": study, "id": de.study_id_for(study)})

        site_name = r["site"]
        if site_name not in site_resolve:
            hit = de.match_site(site_name, matchable_sites, pi_index)
            if hit:
                site_resolve[site_name] = (hit["id"], hit.get("name") or site_name, False)
                match_report["sites_matched"].append(
                    {"from": site_name, "to": hit.get("name"), "id": hit["id"]}
                )
            else:
                site_resolve[site_name] = (de.site_id_for(site_name), site_name, True)
                match_report["sites_new"].append(
                    {"name": site_name, "id": de.site_id_for(site_name)}
                )

    def uniq(items, key):
        seen = set()
        out = []
        for it in items:
            k = it.get(key)
            if k in seen:
                continue
            seen.add(k)
            out.append(it)
        return out

    match_report["studies_matched"] = uniq(match_report["studies_matched"], "id")
    match_report["studies_new"] = uniq(match_report["studies_new"], "id")
    match_report["sites_matched"] = uniq(match_report["sites_matched"], "id")
    match_report["sites_new"] = uniq(match_report["sites_new"], "id")

    by_study = defaultdict(
        lambda: {
            "target_scheduled": 0.0,
            "scheduled": 0.0,
            "screened": 0.0,
            "enrolled": 0.0,
            "sites": set(),
            "pis": set(),
            "dates": [],
            "lplvs": [],
            "canonical": None,
            "is_new": True,
        }
    )
    by_site = defaultdict(
        lambda: {
            "name": None,
            "target_scheduled": 0.0,
            "scheduled": 0.0,
            "screened": 0.0,
            "enrolled": 0.0,
            "studies": set(),
            "pis": set(),
            "nOutcomeRows": 0,
            "is_new": True,
        }
    )

    outcomes = []
    for r in records:
        study_id, study_canon, study_new = study_resolve[r["study"]]
        site_id, site_canon, site_new = site_resolve[r["site"]]
        oid = outcome_id(study_canon, site_canon, r["group"])
        outcomes.append(
            {
                "id": oid,
                "type": "legacyStudySiteOutcome",
                "studyId": study_id,
                "studyName": study_canon,
                "siteId": site_id,
                "siteName": site_canon,
                "group": r["group"],
                "pi": r["pi"],
                "visit1Start": r["visit1_start"],
                "lplv": r["lplv"],
                "targetScheduled": r["target_scheduled"],
                "scheduled": r["scheduled"],
                "screened": r["screened"],
                "enrolled": r["enrolled"],
                "uniqueId": r.get("unique_id"),
                "source": SOURCE,
                "sourceFile": source_file,
                "workbookSiteLabel": r["site_raw"],
                "workbookStudyLabel": r["study"],
                "workbookSheet": r["sheet"],
                "ingestedAt": now,
                "createdAt": now,
                "updatedAt": now,
                "_studyIsNew": study_new,
                "_siteIsNew": site_new,
            }
        )
        agg = by_study[study_id]
        agg["canonical"] = study_canon
        agg["is_new"] = study_new
        for f_src, f_dst in [
            ("target_scheduled", "target_scheduled"),
            ("scheduled", "scheduled"),
            ("screened", "screened"),
            ("enrolled", "enrolled"),
        ]:
            agg[f_dst] += r.get(f_src) or 0
        agg["sites"].add(site_canon)
        if r["pi"]:
            agg["pis"].add(r["pi"])
        if r["visit1_start"]:
            agg["dates"].append(r["visit1_start"])
        if r["lplv"]:
            agg["lplvs"].append(r["lplv"])

        sagg = by_site[site_id]
        sagg["name"] = site_canon
        sagg["is_new"] = site_new
        sagg["nOutcomeRows"] += 1
        sagg["studies"].add(study_canon)
        if r["pi"]:
            sagg["pis"].add(r["pi"])
        for f_src, f_dst in [
            ("target_scheduled", "target_scheduled"),
            ("scheduled", "scheduled"),
            ("screened", "screened"),
            ("enrolled", "enrolled"),
        ]:
            sagg[f_dst] += r.get(f_src) or 0

    studies = []
    for study_id, agg in by_study.items():
        if not agg["is_new"]:
            continue
        studies.append(
            {
                "id": study_id,
                "type": "legacyStudy",
                "name": agg["canonical"],
                "title": agg["canonical"],
                "therapeuticArea": INDICATION,
                "indication": INDICATION,
                "oraProjectNumber": None,
                "sponsor": None,
                "phase": None,
                "status": "Completed",
                "notes": None,
                "source": SOURCE,
                "sourceFile": source_file,
                "editableFields": True,
                "metrics": {
                    "targetScheduled": round(agg["target_scheduled"], 2),
                    "scheduled": round(agg["scheduled"], 2),
                    "screened": round(agg["screened"], 2),
                    "enrolled": round(agg["enrolled"], 2),
                    "nSites": len(agg["sites"]),
                    "nPis": len(agg["pis"]),
                    "nSiteRows": len([o for o in outcomes if o["studyId"] == study_id]),
                    "visit1StartMin": min(agg["dates"]) if agg["dates"] else None,
                    "visit1StartMax": max(agg["dates"]) if agg["dates"] else None,
                    "lplvMin": min(agg["lplvs"]) if agg["lplvs"] else None,
                    "lplvMax": max(agg["lplvs"]) if agg["lplvs"] else None,
                },
                "createdAt": now,
                "updatedAt": now,
                "ingestedAt": now,
            }
        )

    sites = []
    for sid, agg in sorted(by_site.items(), key=lambda x: -(x[1]["enrolled"])):
        if not agg["is_new"]:
            continue
        sites.append(
            {
                "id": sid,
                "type": "legacySite",
                "name": agg["name"],
                "siteCode": de.slugify(agg["name"]).upper().replace("-", "_")[:32],
                "status": "Active",
                "notes": None,
                "relationshipPreference": None,
                "advantages": None,
                "disadvantages": None,
                "relationshipNotes": None,
                "linkedArtemisSiteId": None,
                "source": SOURCE,
                "sourceFile": source_file,
                "editableFields": True,
                "metrics": {
                    "targetScheduled": round(agg["target_scheduled"], 2),
                    "scheduled": round(agg["scheduled"], 2),
                    "screened": round(agg["screened"], 2),
                    "enrolled": round(agg["enrolled"], 2),
                    "nStudies": len(agg["studies"]),
                    "nPis": len(agg["pis"]),
                    "nOutcomeRows": agg["nOutcomeRows"],
                    "studyNames": sorted(agg["studies"]),
                },
                "createdAt": now,
                "updatedAt": now,
                "ingestedAt": now,
            }
        )

    return studies, sites, outcomes, match_report


def upsert_all(studies, sites, outcomes, match_report, dry_run=False):
    de.load_key_from_local_settings()
    if not de.KEY:
        raise SystemExit("COSMOS_KEY missing")

    client = CosmosClient(ENDPOINT, credential=de.KEY)
    db = client.get_database_client(DATABASE_ID)
    existing_studies, existing_sites, existing_outcomes = de.fetch_existing(db)
    by_outcome_id = {o["id"]: o for o in existing_outcomes}

    print("\n=== MATCH REPORT (Old ASO / Enrollment Done) ===")
    print(
        f"Existing Cosmos: {len(existing_studies)} studies, "
        f"{len(existing_sites)} sites, {len(existing_outcomes)} outcomes"
    )
    print(f"Studies matched 1:1: {len(match_report['studies_matched'])}")
    for m in match_report["studies_matched"]:
        print(f"  study  {m['from']!r} -> {m['to']!r} ({m['id']})")
    print(f"Studies NEW: {len(match_report['studies_new'])}")
    for m in match_report["studies_new"]:
        print(f"  + {m['name']!r} ({m['id']})")
    print(f"Sites matched 1:1: {len(match_report['sites_matched'])}")
    for m in match_report["sites_matched"][:40]:
        print(f"  site   {m['from']!r} -> {m['to']!r} ({m['id']})")
    if len(match_report["sites_matched"]) > 40:
        print(f"  ... +{len(match_report['sites_matched']) - 40} more")
    print(f"Sites NEW: {len(match_report['sites_new'])}")
    for m in match_report["sites_new"]:
        print(f"  + {m['name']!r} ({m['id']})")

    to_upsert = []
    for doc in outcomes:
        clean = {k: v for k, v in doc.items() if not k.startswith("_")}
        prev = by_outcome_id.get(clean["id"])
        if prev and prev.get("source") != SOURCE:
            # Shouldn't happen (source in hash) — never overwrite other sources
            match_report["outcomes_existing_preserved"] += 1
            continue
        to_upsert.append(clean)
        match_report["outcomes_new"] += 1

    print(f"\nOutcomes to upsert (this source only): {len(to_upsert)}")
    print(f"New study docs: {len(studies)}")
    print(f"New site docs: {len(sites)}")

    if dry_run:
        print("\nDRY RUN — no writes.")
        if to_upsert:
            print("sample:", json.dumps(to_upsert[0], indent=2)[:700])
        return match_report

    pk_id = PartitionKey(**{"path": "/id"})
    pk_study = PartitionKey(**{"path": "/studyId"})
    db.create_container_if_not_exists(id=STUDIES_CONTAINER, partition_key=pk_id)
    db.create_container_if_not_exists(id=SITES_CONTAINER, partition_key=pk_id)
    db.create_container_if_not_exists(id=OUTCOMES_CONTAINER, partition_key=pk_study)
    studies_c = db.get_container_client(STUDIES_CONTAINER)
    sites_c = db.get_container_client(SITES_CONTAINER)
    outcomes_c = db.get_container_client(OUTCOMES_CONTAINER)

    print(f"Creating {len(studies)} new legacy studies...")
    for doc in studies:
        studies_c.upsert_item(doc)

    matched_ids = {m["id"] for m in match_report["studies_matched"]}
    for st in existing_studies:
        if st["id"] not in matched_ids:
            continue
        full = studies_c.read_item(st["id"], partition_key=st["id"])
        changed = False
        sources = full.get("sources") or []
        if not isinstance(sources, list):
            sources = [sources] if sources else []
        if full.get("source") and full["source"] not in sources:
            sources.append(full["source"])
        if SOURCE not in sources:
            sources.append(SOURCE)
            full["sources"] = sources
            changed = True
        files = full.get("sourceFiles") or ([] if not full.get("sourceFile") else [full["sourceFile"]])
        if not isinstance(files, list):
            files = [files]
        src_path = outcomes[0]["sourceFile"] if outcomes else str(DEFAULT_XLSX)
        if str(src_path) not in files:
            files.append(str(src_path))
            full["sourceFiles"] = files
            changed = True
        if changed:
            full["updatedAt"] = datetime.utcnow().isoformat() + "Z"
            studies_c.upsert_item(full)

    print(f"Creating {len(sites)} new legacy sites...")
    for doc in sites:
        sites_c.upsert_item(doc)

    matched_site_ids = {m["id"] for m in match_report["sites_matched"]}
    for site in existing_sites:
        if site["id"] not in matched_site_ids:
            continue
        full = sites_c.read_item(site["id"], partition_key=site["id"])
        changed = False
        sources = full.get("sources") or []
        if not isinstance(sources, list):
            sources = [sources] if sources else []
        if full.get("source") and full["source"] not in sources:
            sources.append(full["source"])
        if SOURCE not in sources:
            sources.append(SOURCE)
            full["sources"] = sources
            changed = True
        if changed:
            full["updatedAt"] = datetime.utcnow().isoformat() + "Z"
            sites_c.upsert_item(full)

    print(f"Upserting {len(to_upsert)} outcomes...")
    batch = 0
    for doc in to_upsert:
        outcomes_c.upsert_item(doc)
        batch += 1
        if batch % 100 == 0:
            print(f"  ...{batch}/{len(to_upsert)}")
    print("Done.")
    return match_report


def main():
    flags = set(sys.argv[1:])
    dry = "--apply" not in flags
    if "--dry-run" in flags:
        dry = True
    do_relink = "--relink" in flags
    argv = [a for a in sys.argv[1:] if not a.startswith("--")]

    # Optional: reuse dry-eye relink against this source's workbook labels
    if do_relink and not argv:
        print("Relink uses dry-eye helper; prefer matching at ingest time.")
        print("Re-run ingest with --apply after sites/feasibility updates if needed.")
        return

    path = Path(argv[0]) if argv else DEFAULT_XLSX
    if not path.exists():
        raise SystemExit(f"File not found: {path}")

    de.load_key_from_local_settings()
    if not de.KEY:
        raise SystemExit("COSMOS_KEY missing")
    client = CosmosClient(ENDPOINT, credential=de.KEY)
    db = client.get_database_client(DATABASE_ID)

    if "--prune-junk" in flags:
        prune_junk_old_aso(db, dry_run=dry)
        return

    print(f"Parsing Old ASO study tabs from {path} ...")
    records, sheet_stats = parse_workbook(path)
    print(
        f"Parsed {len(records)} site-study rows from "
        f"{sum(1 for x in sheet_stats if x['rows'])} sheets"
    )
    for st in sheet_stats:
        flag = "OK" if st["rows"] else "SKIP"
        print(f"  {flag:4} {st['sheet'][:50]:50} study={st['study']!r:30} rows={st['rows']}")

    if not dry:
        prune_junk_old_aso(db, dry_run=False)

    existing_studies, existing_sites, existing_outcomes = de.fetch_existing(db)

    studies, sites, outcomes, match_report = build_docs(
        records, existing_studies, existing_sites, existing_outcomes, str(path)
    )
    upsert_all(studies, sites, outcomes, match_report, dry_run=dry)
    if dry:
        print("\nRe-run with --apply to write.")


if __name__ == "__main__":
    main()
