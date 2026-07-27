"""
Ingest Anterior Segment Overview.xlsx -> Cosmos (ora-clinical-recruiting / crcscheduling)
Containers ONLY: legacy-studies, legacy-study-site-outcomes
Does NOT write to studies, sites, patients, crcs, events, etc.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

try:
    from azure.cosmos import CosmosClient, PartitionKey
except ImportError:
    print("Install: pip install azure-cosmos openpyxl")
    sys.exit(1)

from openpyxl import load_workbook

DEFAULT_XLSX = Path(r"c:\Users\shue1\Downloads\Anterior Segment Overview.xlsx")
ENDPOINT = os.environ.get("COSMOS_ENDPOINT", "https://ora-clinical-recruiting.documents.azure.com:443/")
KEY = os.environ.get("COSMOS_KEY", "")
DATABASE_ID = os.environ.get("DATABASE_ID", "crcscheduling")
STUDIES_CONTAINER = "legacy-studies"
OUTCOMES_CONTAINER = "legacy-study-site-outcomes"


def load_key_from_local_settings():
    global KEY
    if KEY:
        return
    candidates = [
        Path(r"C:\Users\shue1\OneDrive\Desktop\HoldAll\CHAOS\azure-api-fixed\local.settings.json"),
        Path(__file__).resolve().parents[1] / "api" / "local.settings.json",
    ]
    for p in candidates:
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            vals = data.get("Values") or {}
            KEY = vals.get("COSMOS_KEY") or KEY
            return


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def s(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v).strip()
    if t.startswith("="):  # skip raw formulas if data_only failed
        return None
    return t or None


def slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (name or "").strip().lower()).strip("-")
    return (slug or "unknown")[:80]


def outcome_id(study: str, site: str, group) -> str:
    raw = f"{study}|{site}|{group}"
    h = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]
    return f"legacy-outcome-{h}"


def parse_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Completed Projects"].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and s(r[0]) == "Unique ID")
    col = {
        "unique_id": 0,
        "study": 1,
        "site": 2,
        "group": 3,
        "pi": 4,
        "visit1_start": 5,
        "lplv": 8,
        "target_scheduled": 9,
        "scheduled": 10,
        "screen": 11,
        "enrolled": 13,
    }
    records = []
    for row in rows[hdr_i + 1 :]:
        if not row:
            continue
        study = s(row[col["study"]]) if len(row) > 1 else None
        site = s(row[col["site"]]) if len(row) > 2 else None
        if not study or not site:
            continue
        # skip repeated section headers / blank templates inside the sheet
        if study.lower() in {"study", "unique id", "row labels"}:
            continue
        if site.lower() in {"site", "row labels"}:
            continue
        visit1 = s(row[col["visit1_start"]]) if len(row) > 5 else None
        if visit1 and visit1.lower() in {"visit 1 start", "e002 date"}:
            continue
        records.append(
            {
                "unique_id": s(row[col["unique_id"]]) if len(row) > 0 else None,
                "study": study,
                "site": site,
                "group": num(row[col["group"]]) if len(row) > 3 else None,
                "pi": s(row[col["pi"]]) if len(row) > 4 else None,
                "visit1_start": visit1,
                "lplv": s(row[col["lplv"]]) if len(row) > 8 else None,
                "target_scheduled": num(row[col["target_scheduled"]]) if len(row) > 9 else None,
                "scheduled": num(row[col["scheduled"]]) if len(row) > 10 else None,
                "screened": num(row[col["screen"]]) if len(row) > 11 else None,
                "enrolled": num(row[col["enrolled"]]) if len(row) > 13 else None,
            }
        )

    # Optional indication hints from study tabs
    indication_by_study = {}
    skip = {
        "Scheduled",
        "Screened",
        "Enrolled",
        "Completed Projects",
        "Site Contact",
        "Master Tracker",
        "Template",
    }
    for name in wb.sheetnames:
        if name in skip:
            continue
        sheet_rows = list(wb[name].iter_rows(values_only=True))
        if len(sheet_rows) < 2:
            continue
        vals = list(sheet_rows[1])
        study_name = s(vals[0]) if vals else None
        indication = s(vals[1]) if len(vals) > 1 else None
        if study_name and indication:
            indication_by_study[study_name] = indication
            # also map short tab names
            indication_by_study[name] = indication

    wb.close()
    return records, indication_by_study


def build_docs(records, indication_by_study, source_file: str):
    now = datetime.utcnow().isoformat() + "Z"
    by_study = defaultdict(
        lambda: {
            "target_scheduled": 0.0,
            "scheduled": 0.0,
            "screened": 0.0,
            "enrolled": 0.0,
            "sites": set(),
            "pis": set(),
            "dates": [],
            "lplvs": [],
        }
    )
    outcomes = []
    for r in records:
        study = r["study"]
        study_id = f"legacy-study-{slugify(study)}"
        sid = outcome_id(study, r["site"], r["group"])
        outcomes.append(
            {
                "id": sid,
                "type": "legacyStudySiteOutcome",
                "studyId": study_id,
                "studyName": study,
                "siteName": r["site"],
                "group": r["group"],
                "pi": r["pi"],
                "visit1Start": r["visit1_start"],
                "lplv": r["lplv"],
                "targetScheduled": r["target_scheduled"],
                "scheduled": r["scheduled"],
                "screened": r["screened"],
                "enrolled": r["enrolled"],
                "uniqueId": r["unique_id"],
                "source": "anterior-segment-overview",
                "sourceFile": source_file,
                "ingestedAt": now,
                "createdAt": now,
                "updatedAt": now,
            }
        )
        agg = by_study[study]
        for f_src, f_dst in [
            ("target_scheduled", "target_scheduled"),
            ("scheduled", "scheduled"),
            ("screened", "screened"),
            ("enrolled", "enrolled"),
        ]:
            agg[f_dst] += r.get(f_src) or 0
        agg["sites"].add(r["site"])
        if r["pi"]:
            agg["pis"].add(r["pi"])
        if r["visit1_start"]:
            agg["dates"].append(r["visit1_start"])
        if r["lplv"]:
            agg["lplvs"].append(r["lplv"])

    studies = []
    for study, agg in by_study.items():
        study_id = f"legacy-study-{slugify(study)}"
        indication = indication_by_study.get(study)
        studies.append(
            {
                "id": study_id,
                "type": "legacyStudy",
                "name": study,
                "title": study,
                "therapeuticArea": None,  # editable later in ARTEMIS
                "indication": indication,
                "sponsor": None,
                "phase": None,
                "status": "Completed",
                "notes": None,
                "source": "anterior-segment-overview",
                "sourceFile": source_file,
                "editableFields": True,
                "metrics": {
                    "targetScheduled": round(agg["target_scheduled"], 2),
                    "scheduled": round(agg["scheduled"], 2),
                    "screened": round(agg["screened"], 2),
                    "enrolled": round(agg["enrolled"], 2),
                    "nSites": len(agg["sites"]),
                    "nPis": len(agg["pis"]),
                    "nSiteRows": len([o for o in outcomes if o["studyName"] == study]),
                    "visit1StartMin": min(agg["dates"]) if agg["dates"] else None,
                    "visit1StartMax": max(agg["dates"]) if agg["dates"] else None,
                    "lplvMin": min(agg["lplvs"]) if agg["lplvs"] else None,
                    "lplvMax": max(agg["lplvs"]) if agg["lplvs"] else None,
                },
                "createdAt": now,
                "updatedAt": now,
                "ingestedAt": now,
            }
        )
    return studies, outcomes


def upsert_all(studies, outcomes, dry_run=False):
    load_key_from_local_settings()
    if not KEY:
        raise SystemExit("COSMOS_KEY missing. Set env COSMOS_KEY or use local.settings.json")

    if dry_run:
        print(f"DRY RUN: {len(studies)} studies, {len(outcomes)} outcomes")
        print(json.dumps(studies[:2], indent=2))
        print(json.dumps(outcomes[:2], indent=2))
        return

    client = CosmosClient(ENDPOINT, credential=KEY)
    db = client.get_database_client(DATABASE_ID)
    pk_id = PartitionKey(**{"path": "/id"})
    pk_study = PartitionKey(**{"path": "/studyId"})
    db.create_container_if_not_exists(id=STUDIES_CONTAINER, partition_key=pk_id)
    db.create_container_if_not_exists(id=OUTCOMES_CONTAINER, partition_key=pk_study)
    studies_c = db.get_container_client(STUDIES_CONTAINER)
    outcomes_c = db.get_container_client(OUTCOMES_CONTAINER)

    # Preserve manually edited metadata on re-ingest
    existing_meta = {}
    for item in studies_c.query_items(
        query="SELECT c.id, c.therapeuticArea, c.sponsor, c.phase, c.status, c.notes, c.indication FROM c",
        enable_cross_partition_query=True,
    ):
        existing_meta[item["id"]] = item

    print(f"Upserting {len(studies)} legacy studies...")
    for doc in studies:
        prev = existing_meta.get(doc["id"])
        if prev:
            for k in ("therapeuticArea", "sponsor", "phase", "status", "notes"):
                if prev.get(k) not in (None, ""):
                    doc[k] = prev[k]
            if prev.get("indication") and not doc.get("indication"):
                doc["indication"] = prev["indication"]
            if prev.get("createdAt"):
                doc["createdAt"] = prev["createdAt"]
        studies_c.upsert_item(doc)

    print(f"Upserting {len(outcomes)} site-study outcomes...")
    batch = 0
    for doc in outcomes:
        outcomes_c.upsert_item(doc)
        batch += 1
        if batch % 100 == 0:
            print(f"  ...{batch}/{len(outcomes)}")
    print("Done.")


def main():
    args = [a for a in sys.argv[1:] if a != "--dry-run"]
    dry = "--dry-run" in sys.argv
    path = Path(args[0]) if args else DEFAULT_XLSX
    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    print(f"Parsing {path} ...")
    records, indications = parse_workbook(path)
    print(f"Parsed {len(records)} site-study rows")
    studies, outcomes = build_docs(records, indications, str(path))
    print(f"Built {len(studies)} studies, {len(outcomes)} outcomes")
    upsert_all(studies, outcomes, dry_run=dry)


if __name__ == "__main__":
    main()
