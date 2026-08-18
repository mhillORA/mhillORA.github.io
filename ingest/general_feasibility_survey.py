"""
Ingest Monday.com General Feasibility export -> ARTEMIS site surveys.

Cosmos: ora-clinical-recruiting / crcscheduling
Writes ONLY:
  site-survey-definitions   (one "General Feasibility" template; questions = Excel headers)
  site-survey-assignments   (one submitted assignment per matched ARTEMIS site)
  site-survey-responses     (answers listed on the Sites tab)

Never writes live studies/patients, never touches Study Bid Workbench / bd-budgets.

Usage:
  python ingest/general_feasibility_survey.py
  python ingest/general_feasibility_survey.py --apply
  python ingest/general_feasibility_survey.py "C:\\path\\General_Feasibility.xlsx" --apply
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

try:
    from azure.cosmos import CosmosClient, PartitionKey
except ImportError:
    print("Install: pip install azure-cosmos openpyxl")
    sys.exit(1)

from openpyxl import load_workbook

DEFAULT_XLSX = Path(r"c:\Users\shue1\Downloads\General_Feasibility_1787078071.xlsx")
ALIASES_PATH = Path(__file__).resolve().parent / "general_feasibility_site_aliases.json"
SURVEY_ID = "survey-general-feasibility"
SOURCE = "monday-general-feasibility"

ENDPOINT = os.environ.get("COSMOS_ENDPOINT", "https://ora-clinical-recruiting.documents.azure.com:443/")
KEY = os.environ.get("COSMOS_KEY", "")
DATABASE_ID = os.environ.get("DATABASE_ID", "crcscheduling")

SKIP_HEADERS = {
    "form view",
    "item id (auto generated)",
    "short text",
    "single select",
}

YES_NO = {"yes", "no"}
LEGAL_SUFFIX = re.compile(
    r"\b(llc|inc|incorporated|ltd|limited|pc|p\.c|pa|p\.a|sc|s\.c|pllc|llp|dba)\b",
    re.I,
)


def load_key():
    global KEY, ENDPOINT, DATABASE_ID
    if KEY:
        return
    repo = Path(__file__).resolve().parents[1]
    candidates = [
        Path(r"C:\Users\shue1\OneDrive\Desktop\HoldAll\CHAOS\azure-api-fixed\local.settings.json"),
        repo / "api" / "local.settings.json",
        repo / "data-api-connections.json",
        repo / "data-api-config.json",
    ]
    for p in candidates:
        if not p.exists():
            continue
        data = json.loads(p.read_text(encoding="utf-8"))
        vals = data.get("Values") or {}
        KEY = vals.get("COSMOS_KEY") or KEY
        ENDPOINT = vals.get("COSMOS_ENDPOINT") or ENDPOINT
        DATABASE_ID = vals.get("DATABASE_ID") or DATABASE_ID
        conn = (data.get("cosmosdb-connection") or {}).get("connectionString") or ""
        if conn and not KEY:
            parts = {}
            for chunk in conn.rstrip(";").split(";"):
                if "=" in chunk:
                    k, v = chunk.split("=", 1)
                    parts[k.strip()] = v.strip()
            KEY = parts.get("AccountKey") or KEY
            ENDPOINT = parts.get("AccountEndpoint") or ENDPOINT
            DATABASE_ID = data.get("cosmosdb-connection", {}).get("database") or DATABASE_ID
        if KEY:
            return


def cell(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    t = str(v).strip()
    if not t or t.startswith("="):
        return None
    return t


def slug_header(label: str, idx: int) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (label or "").strip().lower()).strip("-")
    slug = (slug or "q")[:56]
    return f"gf_{idx:02d}_{slug}"


def norm_name(s: str) -> str:
    t = (s or "").strip().lower().replace("&", " and ")
    t = re.sub(r"[^a-z0-9]+", " ", t)
    t = LEGAL_SUFFIX.sub(" ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def norm_email(s: str) -> str:
    return (s or "").strip().lower()


def load_aliases() -> dict[str, str]:
    if not ALIASES_PATH.exists():
        return {}
    raw = json.loads(ALIASES_PATH.read_text(encoding="utf-8"))
    return {str(k): str(v) for k, v in raw.items() if not str(k).startswith("_") and v}


def parse_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if "general feasibility" not in wb.sheetnames:
        raise SystemExit(f"Expected sheet 'general feasibility', found {wb.sheetnames}")
    rows = list(wb["general feasibility"].iter_rows(values_only=True))
    if len(rows) < 4:
        raise SystemExit("Workbook has no data rows")
    headers = [cell(h) or f"Column {i}" for i, h in enumerate(rows[2])]
    records = []
    for row in rows[3:]:
        if not row:
            continue
        site = cell(row[2]) if len(row) > 2 else None
        if not site:
            continue
        values = {}
        for i, header in enumerate(headers):
            values[header] = cell(row[i]) if i < len(row) else None
        records.append(values)
    return headers, records


def infer_type_and_options(header: str, values: list[str | None]):
    filled = [v for v in values if v]
    uniq = sorted(set(filled), key=lambda x: x.lower())
    h = (header or "").lower()
    if "date of response" in h:
        return "date", []
    if h in {"experience (yrs)", "483 year"}:
        return "number", []
    if uniq and all(u.lower() in YES_NO for u in uniq):
        return "select", ["Yes", "No"]
    maxlen = max((len(v) for v in filled), default=0)
    if any(k in h for k in ("equipment", "patient identification", "sponsor", "specialt", "practice setting", "types of ophthalmic")):
        return "textarea", []
    if filled and maxlen > 90:
        return "textarea", []
    if 2 <= len(uniq) <= 8 and maxlen <= 40 and len(filled) >= 8:
        return "select", uniq
    return "text", []


def build_questions(headers, records):
    questions = []
    for i, header in enumerate(headers):
        if header.strip().lower() in SKIP_HEADERS:
            continue
        col_vals = [r.get(header) for r in records]
        qtype, options = infer_type_and_options(header, col_vals)
        questions.append({
            "id": slug_header(header, i),
            "label": header,
            "type": qtype,
            "required": False,
            "options": options,
            "logic": None,
        })
    return questions


def answers_from_record(questions, record):
    out = []
    for q in questions:
        raw = record.get(q["label"])
        out.append({
            "questionId": q["id"],
            "label": q["label"],
            "type": q["type"],
            "skipped": raw is None,
            "value": raw,
        })
    return out


def site_emails(site: dict) -> set[str]:
    keys = [
        "piEmail", "pi2Email", "pi3Email",
        "siteCoordinatorEmail", "siteCoordinator2Email", "siteCoordinator3Email",
    ]
    return {norm_email(site.get(k)) for k in keys if site.get(k)}


def record_emails(record: dict) -> set[str]:
    out = set()
    for k in ("Inv #1 Email", "POC Email"):
        v = norm_email(record.get(k) or "")
        if v:
            out.add(v)
    return out


def name_score(excel_name: str, site: dict) -> float:
    a = norm_name(excel_name)
    candidates = [norm_name(site.get("name") or "")]
    abbr = norm_name(site.get("siteNameAbbreviation") or "")
    if abbr:
        candidates.append(abbr)
    best = 0.0
    for b in candidates:
        if not a or not b:
            continue
        if a == b:
            return 1.0
        if a in b or b in a:
            shorter = min(len(a), len(b))
            if shorter >= 10:
                best = max(best, 0.92)
        ta, tb = set(a.split()), set(b.split())
        if ta and tb:
            jacc = len(ta & tb) / len(ta | tb)
            seq = SequenceMatcher(None, a, b).ratio()
            best = max(best, (jacc * 0.55) + (seq * 0.45))
    return best


def match_record(record, sites, alias_to_site, email_index):
    excel_site = record.get("Site Name") or ""
    emails = record_emails(record)
    for e in emails:
        hit = email_index.get(e)
        if hit:
            return hit, "email", 1.0

    alias_target = alias_to_site.get(excel_site) or alias_to_site.get(norm_name(excel_site))
    if alias_target:
        return alias_target, "alias", 0.99

    scored = []
    for site in sites:
        sc = name_score(excel_site, site)
        if sc >= 0.86:
            scored.append((sc, site))
    scored.sort(key=lambda x: x[0], reverse=True)
    if not scored:
        return None, "unmatched", 0.0
    if len(scored) > 1 and scored[0][0] - scored[1][0] < 0.08:
        return None, "ambiguous", scored[0][0]
    return scored[0][1], "name", scored[0][0]


def fetch_all(container, query="SELECT * FROM c"):
    return list(container.query_items(query=query, enable_cross_partition_query=True))


def assignment_id(site_id: str) -> str:
    return f"asg-gf-{site_id}"[:64]


def response_id(site_id: str) -> str:
    h = hashlib.sha1(site_id.encode("utf-8")).hexdigest()[:12]
    return f"rsp-gf-{h}"


def iso_now():
    return datetime.now(datetime.UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def main():
    ap = argparse.ArgumentParser(description="Match General Feasibility Excel to ARTEMIS sites and load as a site survey")
    ap.add_argument("xlsx", nargs="?", default=str(DEFAULT_XLSX), help="Path to General_Feasibility_*.xlsx")
    ap.add_argument("--apply", action="store_true", help="Write survey definition + matched responses to ARTEMIS Cosmos")
    ap.add_argument("--report", default="", help="Optional path to write match report JSON")
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.exists():
        raise SystemExit(f"Excel not found: {xlsx}")

    load_key()
    if not KEY:
        raise SystemExit("COSMOS_KEY missing. Set env or use data-api-connections.json / local.settings.json")

    headers, records = parse_workbook(xlsx)
    questions = build_questions(headers, records)
    aliases = load_aliases()

    client = CosmosClient(ENDPOINT, credential=KEY)
    db = client.get_database_client(DATABASE_ID)
    sites_c = db.get_container_client("sites")
    sites = fetch_all(sites_c, "SELECT c.id, c.name, c.siteNameAbbreviation, c.pi, c.piEmail, c.pi2Name, c.pi2Email, c.pi3Name, c.pi3Email, c.siteCoordinator, c.siteCoordinatorEmail, c.siteCoordinator2Email, c.siteCoordinator3Email FROM c")

    by_name = {norm_name(s.get("name") or ""): s for s in sites if s.get("name")}
    alias_to_site = {}
    for excel_name, artemis_name in aliases.items():
        site = by_name.get(norm_name(artemis_name))
        if not site:
            print(f"WARN alias target not in ARTEMIS sites: {artemis_name!r}")
            continue
        alias_to_site[excel_name] = site
        alias_to_site[norm_name(excel_name)] = site

    email_index = {}
    for s in sites:
        for e in site_emails(s):
            email_index[e] = s

    matches = []
    unmatched = []
    ambiguous = []
    for rec in records:
        site, how, score = match_record(rec, sites, alias_to_site, email_index)
        row = {
            "excelSite": rec.get("Site Name"),
            "excelName": rec.get("Name"),
            "itemId": rec.get("Item ID (auto generated)"),
            "date": rec.get("Date of Response"),
            "how": how,
            "score": round(score, 3),
            "artemisId": site.get("id") if site else None,
            "artemisName": site.get("name") if site else None,
            "record": rec,
        }
        if how == "unmatched":
            unmatched.append(row)
        elif how == "ambiguous":
            ambiguous.append(row)
        else:
            matches.append(row)

    # One response per ARTEMIS site: keep latest Date of Response
    by_site = defaultdict(list)
    for m in matches:
        by_site[m["artemisId"]].append(m)
    chosen = []
    duplicates = []
    for site_id, rows in by_site.items():
        rows.sort(key=lambda r: str(r.get("date") or ""), reverse=True)
        chosen.append(rows[0])
        for extra in rows[1:]:
            duplicates.append(extra)

    print(f"Excel rows with a site name: {len(records)}")
    print(f"Questions from headers:       {len(questions)}")
    print(f"ARTEMIS live sites:           {len(sites)}")
    print(f"Matched rows:                 {len(matches)}  ({len(chosen)} unique ARTEMIS sites)")
    print(f"Unmatched rows:               {len(unmatched)}")
    print(f"Ambiguous rows:               {len(ambiguous)}")
    print(f"Duplicate excel->same site:   {len(duplicates)}")
    print()
    print("MATCHED")
    for m in sorted(chosen, key=lambda x: (x["artemisName"] or "").lower()):
        print(f"  [{m['how']:5}] {m['excelSite']}  ->  {m['artemisName']}  ({m['artemisId']})")
    if duplicates:
        print()
        print("DUPLICATES (later date kept)")
        for m in duplicates:
            print(f"  skip {m['excelSite']} ({m['date']}) already on {m['artemisName']}")
    if ambiguous:
        print()
        print("AMBIGUOUS (not loaded)")
        for m in ambiguous:
            print(f"  {m['excelSite']}")
    if unmatched:
        print()
        print("UNMATCHED (not loaded - ARTEMIS has no corresponding live site)")
        for m in sorted(unmatched, key=lambda x: (x["excelSite"] or "").lower()):
            print(f"  {m['excelSite']}")

    report = {
        "sourceFile": str(xlsx),
        "surveyId": SURVEY_ID,
        "questionCount": len(questions),
        "excelRows": len(records),
        "matchedSites": [
            {k: v for k, v in m.items() if k != "record"} for m in chosen
        ],
        "unmatched": [{k: v for k, v in m.items() if k != "record"} for m in unmatched],
        "ambiguous": [{k: v for k, v in m.items() if k != "record"} for m in ambiguous],
        "duplicatesSkipped": [{k: v for k, v in m.items() if k != "record"} for m in duplicates],
    }
    if args.report:
        Path(args.report).write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(f"\nWrote report {args.report}")

    if not args.apply:
        print("\nDry run only. Re-run with --apply to write the General Feasibility survey + answers to ARTEMIS.")
        return

    now = iso_now()
    pk = PartitionKey(path="/id")
    defs_c = db.create_container_if_not_exists(id="site-survey-definitions", partition_key=pk)
    asg_c = db.create_container_if_not_exists(id="site-survey-assignments", partition_key=pk)
    rsp_c = db.create_container_if_not_exists(id="site-survey-responses", partition_key=pk)

    existing_def = None
    try:
        existing_def = defs_c.read_item(SURVEY_ID, SURVEY_ID)
    except Exception:
        existing_def = None

    definition = {
        "id": SURVEY_ID,
        "title": "General Feasibility",
        "description": "Imported from the Monday.com General Feasibility form. Questions are the export column headers. Answers are shown on each matched ARTEMIS site.",
        "audience": ["PI", "Coordinator"],
        "status": "active",
        "questions": questions,
        "defaultValues": {},
        "source": SOURCE,
        "sourceFile": xlsx.name,
        "createdAt": (existing_def or {}).get("createdAt") or now,
        "updatedAt": now,
    }
    defs_c.upsert_item(definition)
    print(f"\nUpserted survey definition {SURVEY_ID} ({len(questions)} questions)")

    for m in chosen:
        site_id = m["artemisId"]
        rec = m["record"]
        submitted = rec.get("Date of Response") or now
        if len(str(submitted)) == 10:
            submitted = f"{submitted}T00:00:00Z"
        asg_id = assignment_id(site_id)
        rsp_id = response_id(site_id)
        assignment = {
            "id": asg_id,
            "surveyId": SURVEY_ID,
            "siteId": site_id,
            "targetRole": "pi",
            "targetEmail": rec.get("Inv #1 Email") or rec.get("POC Email") or None,
            "status": "submitted",
            "source": SOURCE,
            "sourceItemId": rec.get("Item ID (auto generated)"),
            "createdAt": now,
            "updatedAt": now,
            "submittedAt": submitted,
        }
        asg_c.upsert_item(assignment)
        response = {
            "id": rsp_id,
            "assignmentId": asg_id,
            "surveyId": SURVEY_ID,
            "siteId": site_id,
            "targetRole": "pi",
            "email": rec.get("Inv #1 Email") or rec.get("POC Email") or None,
            "displayName": rec.get("Name") or rec.get("Investigator #1"),
            "answers": answers_from_record(questions, rec),
            "source": SOURCE,
            "sourceItemId": rec.get("Item ID (auto generated)"),
            "excelSiteName": rec.get("Site Name"),
            "matchHow": m["how"],
            "createdAt": now,
            "updatedAt": now,
            "submittedAt": submitted,
        }
        rsp_c.upsert_item(response)

    print(f"Upserted {len(chosen)} assignments + responses onto matched ARTEMIS sites.")
    print("Open a site on the Sites tab to see General Feasibility answers listed individually.")


if __name__ == "__main__":
    main()
