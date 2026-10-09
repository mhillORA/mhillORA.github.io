"""
Export a survey answer key / codebook:

  - Every question (page, number, label, type)
  - Preset choices for radio / select / multiselect / checkbox
  - Free-text / number / date / address / etc. documented as open response

Usage:
  python ingest/export_survey_answer_key.py survey-cavalry-nk-feasibility
  python ingest/export_survey_answer_key.py survey-cavalry-nk-feasibility --out "C:/Users/shue1/Downloads/Cavalry_answer_key.xlsx"
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from create_cavalry_nk_survey import cosmos_db  # noqa: E402

REPO = Path(__file__).resolve().parents[1]

OPEN_TYPES = {
    "text": "Free text (short)",
    "textarea": "Free text (long)",
    "number": "Number (open)",
    "date": "Date (open)",
    "email": "Free text (email)",
    "rating": "Star rating (1–N)",
}

CHOICE_TYPES = {
    "radio",
    "select",
    "yesno",
    "multiselect",
    "checkboxes",
    "checkbox",
}


def option_label(o) -> str:
    if isinstance(o, dict):
        return str(o.get("label") or o.get("value") or o.get("text") or "").strip()
    return str(o or "").strip()


def response_mode(q: dict) -> str:
    t = str(q.get("type") or "text").lower()
    if t == "address":
        return "Site Address (multipart: Street, City, State, ZIP)"
    if t in OPEN_TYPES:
        return OPEN_TYPES[t]
    if t in ("radio", "yesno"):
        return "Radio — choose one"
    if t == "select":
        return "Dropdown — choose one"
    if t in ("multiselect", "checkboxes", "checkbox"):
        return "Multi-select / checkboxes — choose all that apply"
    if t == "rating":
        return OPEN_TYPES["rating"]
    return f"Other ({t})"


def load_survey(survey_id: str) -> dict:
    db = cosmos_db()
    return db.get_container_client("site-survey-definitions").read_item(survey_id, survey_id)


def page_order(survey: dict) -> dict[str, int]:
    """questionId -> page index."""
    out = {}
    for i, p in enumerate(survey.get("pages") or []):
        for qid in p.get("questionIds") or []:
            out[str(qid)] = i
    return out


def page_title_map(survey: dict) -> dict[str, str]:
    out = {}
    for p in survey.get("pages") or []:
        title = p.get("title") or p.get("key") or ""
        for qid in p.get("questionIds") or []:
            out[str(qid)] = title
    return out


def build_rows(survey: dict) -> list[dict]:
    qmap = {str(q.get("id")): q for q in (survey.get("questions") or []) if q.get("id")}
    p_idx = page_order(survey)
    p_title = page_title_map(survey)
    # Prefer page order; fall back to definition order
    ordered_ids = []
    seen = set()
    for p in survey.get("pages") or []:
        for qid in p.get("questionIds") or []:
            qid = str(qid)
            if qid in qmap and qid not in seen:
                ordered_ids.append(qid)
                seen.add(qid)
    for q in survey.get("questions") or []:
        qid = str(q.get("id") or "")
        if qid and qid not in seen:
            ordered_ids.append(qid)
            seen.add(qid)

    rows = []
    for seq, qid in enumerate(ordered_ids, 1):
        q = qmap[qid]
        t = str(q.get("type") or "text").lower()
        opts = q.get("options") or []
        labels = [option_label(o) for o in opts if option_label(o)]
        logic = q.get("logic") or {}
        show = logic.get("showIf") or {}
        end = logic.get("endSurveyIf") or {}
        branch = ""
        if show:
            branch = f"Show if {show.get('questionId')} = {show.get('equals') or show.get('includes') or ''}"
        if end:
            branch = (branch + "; " if branch else "") + f"End survey if = {end.get('equals') or end.get('includes')}"

        if t in CHOICE_TYPES and labels:
            for i, choice in enumerate(labels, 1):
                rows.append(
                    {
                        "seq": seq,
                        "page": p_title.get(qid) or q.get("category") or "",
                        "pageOrder": p_idx.get(qid, 999) + 1,
                        "questionNum": q.get("docxNum") or "",
                        "questionId": qid,
                        "libraryQuestionId": q.get("libraryQuestionId") or "",
                        "question": q.get("label") or q.get("title") or "",
                        "responseType": response_mode(q),
                        "choiceIndex": i,
                        "availableAnswer": choice,
                        "branchLogic": branch,
                        "required": "Yes" if q.get("required") else "No",
                        "help": q.get("help") or "",
                    }
                )
        elif t == "address":
            for i, part in enumerate(["Street (Line 1 + suite)", "City", "State", "ZIP"], 1):
                rows.append(
                    {
                        "seq": seq,
                        "page": p_title.get(qid) or q.get("category") or "",
                        "pageOrder": p_idx.get(qid, 999) + 1,
                        "questionNum": q.get("docxNum") or "",
                        "questionId": qid,
                        "libraryQuestionId": q.get("libraryQuestionId") or "",
                        "question": q.get("label") or q.get("title") or "",
                        "responseType": response_mode(q),
                        "choiceIndex": i,
                        "availableAnswer": f"[Open field] {part}",
                        "branchLogic": branch,
                        "required": "Yes" if q.get("required") else "No",
                        "help": q.get("help") or "",
                    }
                )
        elif t == "rating":
            max_n = int(q.get("maxStars") or q.get("max") or len(labels) or 5)
            for i in range(1, max_n + 1):
                rows.append(
                    {
                        "seq": seq,
                        "page": p_title.get(qid) or q.get("category") or "",
                        "pageOrder": p_idx.get(qid, 999) + 1,
                        "questionNum": q.get("docxNum") or "",
                        "questionId": qid,
                        "libraryQuestionId": q.get("libraryQuestionId") or "",
                        "question": q.get("label") or q.get("title") or "",
                        "responseType": response_mode(q),
                        "choiceIndex": i,
                        "availableAnswer": str(i),
                        "branchLogic": branch,
                        "required": "Yes" if q.get("required") else "No",
                        "help": q.get("help") or "",
                    }
                )
        else:
            open_label = OPEN_TYPES.get(t) or f"Open ({t})"
            rows.append(
                {
                    "seq": seq,
                    "page": p_title.get(qid) or q.get("category") or "",
                    "pageOrder": p_idx.get(qid, 999) + 1,
                    "questionNum": q.get("docxNum") or "",
                    "questionId": qid,
                    "libraryQuestionId": q.get("libraryQuestionId") or "",
                    "question": q.get("label") or q.get("title") or "",
                    "responseType": response_mode(q),
                    "choiceIndex": 1,
                    "availableAnswer": f"[{open_label}]",
                    "branchLogic": branch,
                    "required": "Yes" if q.get("required") else "No",
                    "help": q.get("help") or "",
                }
            )
    return rows


def write_xlsx(survey: dict, rows: list[dict], out: Path) -> None:
    wb = Workbook()

    # --- Answer key (flat) ---
    ws = wb.active
    ws.title = "Answer Key"
    headers = [
        "Seq",
        "Page",
        "Page #",
        "Q #",
        "Question",
        "Response type",
        "Choice #",
        "Available answer / response mode",
        "Branch logic",
        "Required",
        "Question ID",
        "Library ID",
        "Help",
    ]
    keys = [
        "seq",
        "page",
        "pageOrder",
        "questionNum",
        "question",
        "responseType",
        "choiceIndex",
        "availableAnswer",
        "branchLogic",
        "required",
        "questionId",
        "libraryQuestionId",
        "help",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
    for r_i, row in enumerate(rows, 2):
        for c, k in enumerate(keys, 1):
            ws.cell(r_i, c, row.get(k, ""))
        ws.cell(r_i, 5).alignment = Alignment(wrap_text=True)
        ws.cell(r_i, 8).alignment = Alignment(wrap_text=True)
    widths = [6, 18, 8, 8, 55, 28, 10, 40, 36, 10, 28, 22, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # --- Compact questions summary ---
    ws2 = wb.create_sheet("Questions Summary")
    sum_headers = ["Seq", "Page", "Q #", "Question", "Response type", "Preset answers", "# choices", "Branch logic", "Required"]
    for c, h in enumerate(sum_headers, 1):
        cell = ws2.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font

    # collapse rows by seq
    by_seq: dict[int, list[dict]] = {}
    for row in rows:
        by_seq.setdefault(row["seq"], []).append(row)
    r = 2
    for seq in sorted(by_seq):
        group = by_seq[seq]
        first = group[0]
        answers = [g["availableAnswer"] for g in group]
        preset = " | ".join(answers)
        ws2.cell(r, 1, seq)
        ws2.cell(r, 2, first["page"])
        ws2.cell(r, 3, first["questionNum"])
        ws2.cell(r, 4, first["question"]).alignment = Alignment(wrap_text=True)
        ws2.cell(r, 5, first["responseType"])
        ws2.cell(r, 6, preset).alignment = Alignment(wrap_text=True)
        ws2.cell(r, 7, len(answers))
        ws2.cell(r, 8, first["branchLogic"])
        ws2.cell(r, 9, first["required"])
        r += 1
    for i, w in enumerate([6, 18, 8, 55, 28, 60, 10, 36, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    # --- Cover ---
    ws0 = wb.create_sheet("Cover", 0)
    ws0["A1"] = "ARTEMIS Survey Answer Key"
    ws0["A1"].font = Font(bold=True, size=16)
    ws0["A3"] = "Survey ID"
    ws0["B3"] = survey.get("id")
    ws0["A4"] = "Title"
    ws0["B4"] = survey.get("title")
    ws0["A5"] = "Status"
    ws0["B5"] = survey.get("status")
    ws0["A6"] = "Questions"
    ws0["B6"] = len(by_seq)
    ws0["A7"] = "Answer-key rows"
    ws0["B7"] = len(rows)
    ws0["A8"] = "Exported at (UTC)"
    ws0["B8"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%SZ")
    ws0["A10"] = "How to read"
    ws0["A11"] = (
        "Answer Key lists one row per preset choice (radio / dropdown / multi-select). "
        "Open responses (free text, number, date) show a single row describing the response mode. "
        "Site Address is one question with four open fields (Street, City, State, ZIP). "
        "Questions Summary collapses choices onto one row per question."
    )
    ws0["A11"].alignment = Alignment(wrap_text=True)
    ws0.merge_cells("A11:B14")
    ws0.column_dimensions["A"].width = 22
    ws0.column_dimensions["B"].width = 80

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser(description="Export survey answer key (preset choices + open types)")
    ap.add_argument("survey_id", nargs="?", default="survey-cavalry-nk-feasibility")
    ap.add_argument("--out", help="Output .xlsx path")
    args = ap.parse_args()

    survey = load_survey(args.survey_id)
    rows = build_rows(survey)
    title_slug = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (survey.get("title") or args.survey_id))[:60]
    out = Path(args.out) if args.out else Path.home() / "Downloads" / f"{title_slug}_answer_key.xlsx"
    write_xlsx(survey, rows, out)

    n_q = len({r["seq"] for r in rows})
    n_choice = sum(1 for r in rows if not str(r["availableAnswer"]).startswith("["))
    print(f"Survey: {survey.get('id')} — {survey.get('title')}")
    print(f"Questions: {n_q}")
    print(f"Answer-key rows: {len(rows)} (preset choice rows~{n_choice})")
    print(f"Wrote: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
